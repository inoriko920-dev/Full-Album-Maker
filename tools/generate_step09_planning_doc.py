from __future__ import annotations
import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
DOC_ROOT=ROOT/"docs"/"beat-animation"
OUT=DOC_ROOT/"STEP09_BEAT_ANIMATION_EXPANSION_EDITOR_CONTROLS_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOC_ROOT/"SOURCE_OF_TRUTH.md"

SECTIONS=[
("1. Tujuan Utama",[
"Beat Animation dapat dipilih dari Property Inspector tanpa edit JSON.",
"Final-render adapter diperluas ke Background, Song Visual, Spectrum, serta subset Text/Title yang sungguh didukung FFmpeg.",
"Intensity per layer 0-200% tanpa mengubah timing/event detection.",
"Setiap perubahan UI harus Undo/Redo-safe.",
"Assignment STEP08 tanpa intensity tetap dibaca sebagai 100%."]),
("2. Prinsip Non-Negotiable",[
"UI hanya menawarkan preset yang punya jalur render yang dinyatakan supported.",
"Beat math tetap berasal dari STEP04-07; Property Inspector tidak menghitung beat dan renderer tidak menjalankan librosa.",
"Intensity hanya mengubah amplitude visual binding.",
"Stable main tidak di-merge pada STEP09."]),
("3. Assignment Schema",[
'layer.animation["beat_v1"] menyimpan enabled, presets, intensity.',
"Intensity optional; default 1.0; range 0.0..2.0.",
"Nama key tetap beat_v1 agar tidak perlu migration."]),
("4. Semantik Intensity",[
"effective binding amount = preset amount * intensity.",
"Intensity 0 menyimpan assignment tetapi visual neutral; 1 default; 2 maksimum UI.",
"Safe bounds VisualPropertyState tetap berlaku."]),
("5. Layer Capability Matrix",[
"song_cover/vinyl: FULL untuk semua 6 preset.",
"background/song_visual/spectrum: FULL; Rotation hanya jika pivot/base rotation aman.",
"text/song_title: FULL subset tanpa Rotation Nudge.",
"Capability API adalah source of truth untuk UI, preflight, dan renderer."]),
("6. Preset Labels",[
"subtle_beat_pulse=Beat Pulse; bass_pulse=Bass Pulse; strong_punch=Strong Punch.",
"onset_flash=Onset Flash; rotation_nudge=Rotation Nudge; energy_breathe=Energy Breathe."]),
("7. Property Inspector",[
"Section BEAT ANIMATION: Aktif, Preset, Intensity 0-200%, status render.",
"Section hanya visible pada supported layer; locked layer controls disabled.",
"Multi-preset lama ditampilkan Custom/Multi sampai user memilih preset baru."]),
("8. Undo/Redo Command",[
"Tambahkan SetLayerAnimationValue untuk nested layer.animation key.",
"Command menyimpan old value + missing state dan menghasilkan inverse command.",
"UI dilarang memodifikasi layer.animation langsung."]),
("9. Session API",[
"EditorSession.set_beat_animation(layer_id, enabled, preset, intensity).",
"enabled=false menghapus beat_v1 agar project bersih.",
"UI V1 memilih satu preset; parser tetap menerima multi-preset lama."]),
("10. Assignment Parser",[
"BeatAnimationAssignment mendapat intensity float default 1.0.",
"binding_set meng-clone bindings dengan amount dikalikan intensity.",
"Intensity tidak mengubah signal timestamps atau analysis cache identity."]),
("11. Capability API",[
"BeatLayerCapability: supported_presets, final_render_level, reason.",
"beat_capability_for_layer dan preset_supported_for_layer dipakai bersama UI/preflight/tests."]),
("12. Background Adapter",[
"Beat suffix ditempatkan setelah source RGBA dan sebelum overlay.",
"Existing zoom/pan/background motion tetap berjalan.",
"Solid/effect/image/video path harus covered.",
"Pivot-safe overlay expressions dipakai untuk dynamic scale."]),
("13. Song Visual Adapter",[
"Beat sendcmd ditempatkan setelah PTS digeser ke project-global time.",
"Control window di-intersect dengan song visual interval.",
"video_speed, transition, loop/freeze tidak diubah."]),
("14. Spectrum Adapter",[
"showfreqs/showwaves/circular tetap digerakkan audio asli.",
"Beat V2 hanya memodulasi hasil visual setelah visualizer.",
"Reactive scale, bands, smoothing, audio master routing tidak berubah."]),
("15. Text/Song Title Adapter",[
"Gunakan named drawtext + sendcmd.",
"Scale/Zoom -> fontsize; X/Y -> x/y; glow -> borderw; opacity -> alpha.",
"Rotation runtime tidak didukung pada STEP09, sehingga Rotation Nudge disembunyikan/ditolak untuk text/title."]),
("16. Text Command File",[
"Timestamp sendcmd menggunakan project time.",
"Target drawtext@beat_text_<token>.",
"Neutral reset mengembalikan fontsize/borderw ke base."]),
("17. Accurate Preview",[
"apply_beat_snapshot diperluas dengan ephemeral _beat_snapshot_font_scale dan _beat_snapshot_glow.",
"Marker hanya pada clone preview dan tidak dipersist.",
"Final render tetap memakai sendcmd runtime."]),
("18. Fast Preview Text/Title",[
"PreviewCanvas memakai effective geometry dan font scale dari runtime.",
"Selection box boleh mengikuti effective geometry, tetapi transform gesture selalu menulis base Transform.",
"Glow fast-preview hanya approximation; pixel authority tetap Accurate Preview."]),
("19. Editor Status & UX",[
"Status assignment, analyzing, runtime-ready, unsupported preset, render-limited harus jelas.",
"Tidak ada modal dialog untuk perubahan normal."]),
("20. Backward Compatibility",[
"Legacy assignment tanpa intensity = 1.0.",
"Multi-preset tetap bisa dibaca dan dirender.",
"No beat assignment -> no analyzer/no extra filter.",
"Spectrum tanpa Beat V2 tetap graph lama."]),
("21. Render Preflight Safety",[
"Beat layers >4 fail; unsupported preset/layer fail; text Rotation fail/hidden.",
"Non-center rotation yang tak aman fail sebelum FFmpeg.",
"Missing analysis disiapkan di luar compiler; command rows >250k fail."]),
("22. Performance",[
"Runtime rebuild async; intensity change memakai cache analysis yang ada.",
"Command generation hanya active envelope windows dan <=30Hz.",
"Tidak membuat frame-by-frame project data."]),
("23. Module/File Plan",[
"Modify assignment, render control, runtime, render_graph/v13, preview, inspector, editor commands/session/workspace.",
"Helpers: beat_layer_capabilities.py, beat_text_render_control.py.",
"Focused tests + real FFmpeg smoke."]),
("24. Test Matrix",[
"Minimal 48 cases: legacy/intensity, Undo/Redo, UI capability filtering, background/song_visual/spectrum/text render, preview parity, cache reuse, regressions STEP04-08, real FFmpeg smokes."]),
("25. CI Gate",[
"Ubuntu+Windows focused tests and regressions.",
"Ubuntu real FFmpeg: background, song_visual, spectrum, drawtext.",
"Qt offscreen Inspector smoke; stable main SHA verify."]),
("26. Urutan Implementasi",[
"DOCX + SOURCE_OF_TRUTH first; assignment/capability; command/session; inspector UI; background; song visual; spectrum; text/title; preview; tests/CI."]),
("27. Acceptance Gate",[
"G1 source of truth; G2 assignment; G3 Undo; G4 UI; G5 background; G6 song visual; G7 spectrum; G8 text/title; G9 no fake option; G10 backward; G11 regression; G12 real FFmpeg; G13 stable main."]),
("28. Definition of Done",[
"User dapat memilih Beat Animation + intensity dari Inspector.",
"Changes persist and Undo/Redo.",
"Background/Song Visual/Spectrum final render Beat modulation.",
"Text/Title supported subset jujur dan teruji.",
"No analyzer in compiler; no duplicated beat math."]),
("29. Handoff STEP10",[
"STEP10 boleh memperluas katalog, genre presets, BPM-sync Vinyl, Spectrum V2 beat punch/bass boost, AI Agent safe actions.",
"Fondasi STEP04-09 tidak diubah tanpa keputusan arsitektur terdokumentasi."]),
("30. Referensi Teknis",[
"FFmpeg drawtext commands mendukung reinit serta x, y, alpha, fontsize, fontcolor, bordercolor, borderw; dipakai untuk Text/Title Beat V2 subset."]),
]

def build():
    DOC_ROOT.mkdir(parents=True,exist_ok=True)
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(9.5)
    for name,size,color in [("Title",22,"17365D"),("Heading 1",15,"17365D"),("Heading 2",12,"2F5597")]:
        st=doc.styles[name]; st.font.name="Aptos Display"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 09\nBeat Animation Expansion & Editor Controls")
    doc.add_paragraph("Branch: feature/beat-animation-engine-v2\nStable main: 584c94774e6197ecf60ecedb3bffd5a8797e7737\nSTEP08 verified head: 26eec01711e6f4db5f6df528f3c5b1ddef2b7cff\nTanggal: 7 Oktober 2026 (WIB)")
    doc.add_heading("Arsitektur Inti",level=1)
    doc.add_paragraph("STEP09 tidak mengubah beat math. Ia menambah capability-aware editor controls dan renderer adapters di atas STEP04-08.")
    for title,items in SECTIONS:
        doc.add_heading(title,level=1)
        for item in items:
            doc.add_paragraph(item,style="List Bullet")
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 09 • 7 Oct 2026")
    doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8") if SOT.exists() else "# Beat Animation Engine V2 — Source of Truth\n"
    entry=f"- `{OUT.name}` — SHA256 `{digest}`"
    lines=[line for line in text.splitlines() if OUT.name not in line]
    pos=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(pos,entry); lines.insert(pos+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__": build()
