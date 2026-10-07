from __future__ import annotations
from pathlib import Path
from hashlib import sha256
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/"docs"/"beat-animation"
OUT=DOCS/"STEP14_FINAL_PERFORMANCE_RELIABILITY_HARDENING_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOCS/"SOURCE_OF_TRUTH.md"

SECTIONS = [
("1. Audit / Release Blockers", [
"pyproject masih requires-python >=3.11 sementara librosa 1.0.0 membutuhkan Python >=3.12; metadata harus diselaraskan.",
"build/requirements-windows.lock belum memuat scientific Beat Analysis stack; final portable build harus mem-pin numpy/scipy/librosa beserta runtime transitive yang sudah terbukti di Windows CI.",
"AccuratePreviewService membangun Beat runtime berulang; perlu memo satu-runtime dengan invalidation berdasarkan document signature + source size/mtime.",
"External cancellation saat ensure_beat_analysis menunggu worker harus dipropagasi ke AudioAnalysisService.cancel().",
"Cache corruption harus tetap disposable: quarantine -> miss -> re-analysis; quarantine/temp harus bounded."
]),
("2. Cancellation Contract", [
"AudioAnalysisService.wait_for_idle menerima optional cancel_event dan polling interval tanpa mematahkan API lama.",
"Cancellation checkpoints: FFT block, sebelum/sesudah beat_track, sesudah onset_detect, sebelum compression, sebelum curves replace, sebelum manifest commit marker.",
"Cancel tidak menghasilkan result_ready atau request_failed error toast; temporary PCM dibersihkan."
]),
("3. Cache Maintenance & Recovery", [
"Tambah CacheMaintenanceReport: removed_partial_entries, removed_temp_files, removed_quarantine_entries, reclaimed_bytes.",
"maintenance defaults: stale partial 3600 s, max quarantine 32 entries, max quarantine 256 MiB.",
"Valid cache tidak dihapus; corrupt manifest/curves otomatis quarantine dan ensure_beat_analysis membangun replacement fresh."
]),
("4. Accurate Preview Runtime Memo", [
"Memo key = document content signature + tuples asset_id/resolved source/size/mtime_ns.",
"Same document/source reuse runtime; document edit, source size change, mtime change, missing source invalidate.",
"Expose clear_beat_runtime_cache(); memo hanya satu runtime terakhir agar bounded."
]),
("5. Hash Work Reduction", [
"ensure_beat_analysis menyimpan first-pass cache hits.",
"Setelah job selesai gunakan AudioAnalysisService.last_result(asset_id) sebelum fallback peek_cached, sehingga file besar tidak langsung di-SHA256 kedua kali.",
"Content identity tetap SHA-256; integrity tidak diturunkan."
]),
("6. Python / Dependency Contract", [
'pyproject requires-python menjadi ">=3.12".',
"Runtime deps declare PySide6, numpy>=2.5,<3, scipy>=1.18,<2, librosa==1.0.0.",
"Windows lock memasukkan exact scientific runtime pins yang sudah ter-resolve pada Windows CI STEP13.",
"PyInstaller tetap 6.22.3; onedir dipertahankan."
]),
("7. Windows Lock Pins", [
"numpy 2.5.3; scipy 1.18.1; librosa 1.0.0; numba 0.68.0; llvmlite 0.50.0; scikit-learn 1.9.1.",
"joblib 1.6.0; soundfile 0.14.0; soxr 1.1.0; pooch 1.9.0; lazy_loader 0.6; msgpack 1.2.3.",
"cffi 2.1.1; pycparser 3.0; cloudpickle 3.1.2; threadpoolctl 3.7.0; narwhals 2.26.0.",
"requests 2.34.2; urllib3 2.8.0; certifi 2026.7.22; charset_normalizer 3.5.2; idna 3.20; typing-extensions 4.16.0; platformdirs 4.11.14; decorator 5.3.1."
]),
("8. Frozen Beat Smoke V2", [
"Windows PyInstaller smoke imports numpy/scipy/librosa/numba dan Beat modules.",
"Generate synthetic 3s 120-BPM float32 PCM; run analyze_pcm; tempo octave error <5 BPM.",
"Smoke harus jalan tanpa global Python/PYTHONPATH dan print STEP14_FROZEN_BEAT_OK."
]),
("9. Long Project / Memory Strategy", [
"Heavy analysis worker tetap 1; PCM tetap disk memmap; FFT block 1024 frames.",
"Accurate preview memo bounded 1 runtime; Beat render layer cap 4; command batches/ops/bytes guards dari STEP13 dipertahankan.",
"Stress: 2h event program random seek, 200-song/3h project compile/estimator, 1000 accurate-preview ticks memo behavior."
]),
("10. Error / Recovery", [
"SOURCE_MISSING/SOURCE_CHANGED actionable; CACHE corrupt transparently reanalyze; CACHE_WRITE_FAILED retryable.",
"CANCELLED silent cancellation state; ANALYSIS_TIMEOUT actionable retry/cancel.",
"Frozen missing dependency adalah STEP14/15 build gate failure."
]),
("11. Capability / Notice Audit", [
"CAPABILITIES.json declare scientific runtime, Beat Analysis, Beat Animation V2, Vinyl BPM Sync, AI Beat actions.",
"THIRD_PARTY_NOTICES audit includes bundled scientific dependencies; librosa ISC and binary dependency notices reviewed."
]),
("12. Test Matrix Highlights", [
"Python>=3.12 metadata; Windows exact lock; no conflicting pins.",
"Cache maintenance/recovery/corruption replacement; cancellation before commit marker.",
"Service external cancel; last_result/invalidate; avoid second hash.",
"Accurate preview memo hit/invalidation/clear.",
"2h event seek deterministic; 200-song/3h stress bounded.",
"Windows/Ubuntu scientific imports; Windows frozen synthetic Beat analysis.",
"Full STEP04-13 regression."
]),
("13. Acceptance Gates", [
"G1 STEP14 DOCX committed before code.",
"G2 Python/librosa/scientific declarations consistent.",
"G3 Windows frozen Beat synthetic analysis PASS.",
"G4 External cancellation responsive; no stale callback/valid commit after cancellation checkpoint.",
"G5 Corrupt cache auto-reanalyzed; quarantine/temp bounded.",
"G6 Same doc/audio accurate preview reuses runtime; changes invalidate.",
"G7 Completed analysis avoids immediate unnecessary rehash.",
"G8 Long-project synthetic gates bounded.",
"G9 STEP04-13 regression PASS Windows/Ubuntu.",
"G10 Feature branch Windows onedir imports scientific stack.",
"G11 stable main remains v1.5.0 baseline."
]),
("14. Handoff to STEP15", [
"STEP15 only final Windows portable/release-candidate build: versioning, clean build, full tests, extracted ZIP smoke without global Python/FFmpeg/API key, Beat frozen runtime smoke, real export, checksums, release notes, GO/NO-GO.",
"No new feature work after STEP14 except test-proven bug fixes."
]),
]

def add_bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after=Pt(2)
    p.add_run(text)

def build():
    DOCS.mkdir(parents=True,exist_ok=True)
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(9.5)
    title=doc.add_paragraph()
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=title.add_run("STEP 14\nFinal Performance & Reliability Hardening")
    r.bold=True; r.font.size=Pt(22); r.font.color.rgb=RGBColor(23,54,93)
    p=doc.add_paragraph()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Full-Album-Maker • Beat Animation Engine V2 • Pre-Release Source of Truth").bold=True
    doc.add_paragraph("Repository: inoriko920-dev/Full-Album-Maker\nBranch: feature/beat-animation-engine-v2\nStable main baseline: 584c94774e6197ecf60ecedb3bffd5a8797e7737\nSTEP13 verified head: 945bc25a527fe4ee9c3c278ce31b1649bfa6830b\nDate: 7 Oct 2026 WIB")
    doc.add_page_break()
    for heading,bullets in SECTIONS:
        doc.add_heading(heading,level=1)
        for item in bullets: add_bullet(doc,item)
    doc.save(OUT)
    digest=sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    tick=chr(96)
    entry=f"- {tick}{OUT.name}{tick} — SHA256 {tick}{digest}{tick}"
    lines=[line for line in text.splitlines() if OUT.name not in line]
    idx=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(idx,entry); lines.insert(idx+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"{OUT} SHA256={digest}")

if __name__=="__main__":
    build()
