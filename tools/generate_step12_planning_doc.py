from __future__ import annotations
import hashlib
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCROOT=ROOT/"docs"/"beat-animation"
OUT=DOCROOT/"STEP12_EVENT_PHASE_MODULATOR_ADVANCED_MOTION_EFFECTS_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOCROOT/"SOURCE_OF_TRUTH.md"
DARK="17365D"; BLUE="D9EAF7"; GREEN="E2F0D9"; YELLOW="FFF2CC"

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)
def margins(cell,v=46,h=54):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for n,val in [("top",v),("bottom",v),("start",h),("end",h)]:
        node=tcMar.find(qn(f"w:{n}"))
        if node is None: node=OxmlElement(f"w:{n}"); tcMar.append(node)
        node.set(qn("w:w"),str(val)); node.set(qn("w:type"),"dxa")
def table(doc,headers,rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,DARK); margins(c)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(7.8)
    for rowv in rows:
        row=t.add_row()
        for i,v in enumerate(rowv):
            row.cells[i].text=str(v); margins(row.cells[i])
            for p in row.cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(7.5)
    return t
def callout(doc,title,text,fill=BLUE):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,fill); margins(c,70,85)
    p=c.paragraphs[0]; r=p.add_run(title+" — "); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text)
def code(doc,text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,"F7F7F7"); margins(c,55,75)
    r=c.paragraphs[0].add_run(text); r.font.name="Consolas"; r.font.size=Pt(7.2)
def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(1.5); p.add_run(text)

def build():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.56); sec.bottom_margin=Inches(.56); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(8.9)
    for n,s,c in [("Title",22.5,"17365D"),("Heading 1",14.2,"17365D"),("Heading 2",11.0,"2F5597")]:
        doc.styles[n].font.name="Aptos Display"; doc.styles[n].font.size=Pt(s); doc.styles[n].font.color.rgb=RGBColor.from_string(c)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 12\nEvent Phase Modulator & Advanced Motion Effects")
    callout(doc,"Status STEP","Signed/event-phase motion di atas signal STEP06: alternating wobble, bass sway, deterministic camera shake, beat bounce, four-way kick, dan Spark Burst V1. Semua efek seek-deterministic dan preview/render parity-safe.",GREEN)
    table(doc,["Item","Nilai"],[
        ("Repository","inoriko920-dev/Full-Album-Maker"),
        ("Development branch","feature/beat-animation-engine-v2"),
        ("Stable main","584c94774e6197ecf60ecedb3bffd5a8797e7737 (v1.5.0)"),
        ("STEP11 head","19dcaaeff4f976f357a1ff177d300426b30d8c91"),
        ("Target","6 motion presets + bounded deterministic particles"),
    ])
    doc.add_page_break()
    doc.add_heading("1. Architecture Boundary",1)
    for x in [
        "STEP04 AudioAnalysisResult tidak berubah.",
        "STEP05 MusicEventTimeline tidak berubah.",
        "STEP06 SignalTrigger/envelope tetap source of truth timing.",
        "STEP12 menambahkan ordinal, signed direction, deterministic noise dan particle identity.",
        "Tidak ada random global, beat detector baru, atau mutable playback state.",
    ]: bullet(doc,x)
    code(doc,'layer.animation["beat_v1"] = {\n  "enabled": true,\n  "presets": ["club_punch"],\n  "intensity": 1.0,\n  "motion_preset": "camera_shake",\n  "motion_intensity": 0.8\n}')
    doc.add_heading("2. EventPhaseEngine",1)
    code(doc,'PhaseSample:\n  channel, event_tick, event_index\n  local_progress_0_1\n  alternating_sign (-1/+1)\n  cardinal_index (0..3)\n  noise_x/noise_y/noise_rotation (-1..1)')
    callout(doc,"Determinism","event_index berasal dari trigger terurut. Shake noise berasal dari SHA-256 trigger identity + 30Hz time bucket, lalu smooth interpolation.",YELLOW)
    doc.add_heading("3. Advanced Motion Presets",1)
    table(doc,["Preset","Source","Output","Default intensity"],[
        ("alternating_wobble","STRONG_BEAT","rotation ±3°","0.85"),
        ("bass_sway","BASS","x ±0.030","0.90"),
        ("camera_shake","STRONG_BEAT","x/y noise + rotation noise","0.80"),
        ("beat_bounce","BEAT","y -0.030","0.75"),
        ("four_way_kick","STRONG_BEAT","left/up/right/down","0.90"),
        ("spark_burst","STRONG_BEAT","6 micro-sparks","0.85"),
    ])
    doc.add_heading("4. Motion Math",1)
    code(doc,'alternating_wobble: rotation = strong_signal * alternating_sign * 3deg * intensity\nbass_sway: x = bass_signal * alternating_sign * 0.030 * intensity\ncamera_shake: x=noise_x*0.014*m; y=noise_y*0.011*m; rot=noise_rotation*0.80deg*m\nbeat_bounce: y=-beat_signal*0.030*intensity\nfour_way_kick: vector[event_index%4] * strong_signal * 0.028 * intensity')
    doc.add_heading("5. Spark Burst V1",1)
    code(doc,'ParticleBurst:\n  trigger=STRONG_BEAT\n  lifetime=220ms\n  particles=6\nParticle:\n  deterministic angle/radius/travel/size/brightness/seed\nprogress=(tick-event_tick)/lifetime\nradius=start+travel*ease_out(progress)\nalpha=(1-progress)^2')
    for x in [
        "Fast Preview menggambar six micro-sparks relatif terhadap effective layer center.",
        "Accurate Preview menyimpan marker particles hanya pada clone frame.",
        "Final render memakai 6 named drawbox slots per Spark layer dan sendcmd runtime x/y/w/h/color.",
        "Spark-enabled layers maksimum 2; command rows tetap mengikuti cap 250000.",
    ]: bullet(doc,x)
    doc.add_heading("6. UI & Compatibility",1)
    table(doc,["Layer","Motion V1"],[
        ("song_cover","Semua 6"),("vinyl","Semua 6"),("background","Semua 6"),
        ("song_visual","Semua 6"),("spectrum","Semua 6"),
        ("text/song_title","Beat Bounce"),
    ])
    code(doc,'BEAT ANIMATION\n  Aktif [x]\n  Preset Beat [Club Punch]\n  Intensity [110%]\n  Motion [Camera Shake]\n  Motion Intensity [80%]')
    doc.add_heading("7. Render Integration",1)
    for x in [
        "BeatVisualRuntime merges VisualPropertyState + AdvancedMotionState before apply_visual_state.",
        "beat_render_control accepts X/Y offsets and emits overlay@id runtime x/y commands.",
        "Scale/rotation/overlay/glow/spark commands share project-global timestamps.",
        "Pivot compensation remains based on overlay_w/overlay_h.",
        "drawbox runtime commands provide bounded Spark slots.",
    ]: bullet(doc,x)
    doc.add_heading("8. Performance & Safety",1)
    table(doc,["Guard","V1"],[
        ("Motion lookup","O(log N + local overlap)"),("Shake","No full-frame buffer"),
        ("Particles","6 per burst"),("Spark layers","max 2"),("Control rate","<=30Hz"),
        ("X/Y clamp","±0.15 normalized"),("Rotation clamp","±12°"),("Command rows","<=250000"),
    ])
    doc.add_page_break()
    doc.add_heading("9. Module Plan",1)
    code(doc,'new:\n  event_phase_modulator.py\n  advanced_motion_contract.py\n  advanced_motion_engine.py\n  spark_burst_engine.py\n  spark_render_control.py\nmodify:\n  beat_animation_assignment.py\n  beat_visual_runtime.py\n  beat_render_control.py\n  render_graph.py\n  preview_scene.py\n  property_inspector.py\n  editor_session.py')
    doc.add_heading("10. Test Matrix",1)
    tests=[
        "Phase ordinal stable","Alternating sign flips per event","Four-way cardinal cycle","Same-tick seek deterministic",
        "Input order invariant","Overlapping trigger deterministic winner","Noise bounded","Noise interpolated",
        "No global RNG dependency","Legacy assignment motion none","Motion intensity 0..2 validation",
        "Alternating Wobble exact","Bass Sway exact","Camera Shake deterministic","Shake neutral at signal zero",
        "Beat Bounce exact","Four-Way exact","Motion clamp","Visual+motion combination","Undo/Redo motion",
        "Inspector legacy load","Inspector capability filter","Spark 6 particles","Spark lifetime ends",
        "Spark stable same event","Spark seed changes next event","Spark layer guard","Snapshot motion baked",
        "Snapshot particle clone-only","Fast preview effective motion","No double spark accurate frame",
        "Render x/y accepted","Overlay command deterministic","Pivot preserved","Spark drawbox bounded",
        "Spark neutral reset","Command cap","Real FFmpeg sway","Real FFmpeg shake","Real FFmpeg spark",
        "STEP11 regression","STEP10 regression","STEP09 regression","STEP08 regression","STEP07 regression",
        "STEP06 regression","STEP05 regression","STEP04 regression",
    ]
    table(doc,["ID","Test"],[(f"{i:02d}",t) for i,t in enumerate(tests,1)])
    doc.add_heading("11. Acceptance Gate",1)
    table(doc,["Gate","PASS Criteria"],[
        ("G1 Source of Truth","STEP12 DOCX committed before production code"),
        ("G2 Phase","Ordinal/sign/direction deterministic"),
        ("G3 Shake","Seek-safe deterministic noise"),
        ("G4 Motion","5 geometry motion effects preview/render parity"),
        ("G5 Spark","6-particle bounded preview + final render"),
        ("G6 UI","Motion preset/intensity capability-filtered"),
        ("G7 Undo","Motion assignment Undo/Redo safe"),
        ("G8 Render","Dynamic overlay x/y runtime commands PASS"),
        ("G9 Safety","Property clamps enforced"),
        ("G10 Performance","No full-album per-frame cache"),
        ("G11 Regression","STEP04-11 PASS Windows/Ubuntu"),
        ("G12 Real FFmpeg","Sway/shake/spark smoke PASS"),
        ("G13 Stable Main","main remains v1.5.0 baseline"),
    ])
    callout(doc,"Gate decision","Semua G1-G13 wajib PASS sebelum STEP12 ditutup.",GREEN)
    doc.add_heading("12. Handoff STEP13",1)
    doc.add_paragraph("STEP13 dapat menambah AI action untuk motion/Spark, combo visual+motion presets, dan performance hardening/command compaction.")
    doc.add_heading("13. Technical Reference",1)
    doc.add_paragraph("FFmpeg overlay mendukung runtime x/y commands; drawbox mendukung commands yang sama dengan options (x/y/w/h/color/thickness).")
    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 12 • 7 Oct 2026"); r.font.size=Pt(7); r.font.color.rgb=RGBColor(105,105,105)
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    entry=f"- `{OUT.name}` — SHA256 `{digest}`"
    lines=[line for line in text.splitlines() if OUT.name not in line]
    idx=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(idx,entry); lines.insert(idx+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__":
    build()
