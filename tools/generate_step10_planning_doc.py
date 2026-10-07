from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / "docs" / "beat-animation"
OUT = DOC_ROOT / "STEP10_ADVANCED_BEAT_ANIMATION_CATALOG_MUSIC_PRESETS_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT = DOC_ROOT / "SOURCE_OF_TRUTH.md"

SECTIONS = [
("1. Tujuan STEP 10", [
"Mengembangkan katalog Beat Animation dari 6 menjadi 18 preset tanpa membuat engine baru per efek.",
"Menyediakan 9 Music Style project recipes, metadata registry tunggal, dan Vinyl BPM Sync berbasis TempoSummary STEP04.",
"Menjaga analyzer, MusicEventTimeline, AnimationSignalEngine, VisualBindingEngine, preview/render parity, serta stable main tetap tidak berubah."
]),
("2. Batasan", [
"Masuk: 18 preset berbasis channel STEP06, registry metadata, 9 Music Style, Apply Music Style, BPM-sync controls pada Vinyl, preview/render BPM Sync.",
"Ditunda: random camera shake, particles, DROP/SECTION_CHANGE, kick/snare classifier, projectM/GPU shader, AI Agent natural-language execution, dan klaim 100 efek."
]),
("3. Arsitektur", [
"AudioAnalysisResult/TempoSummary memberi data tempo. MusicEventTimeline -> AnimationSignalEngine -> BeatPresetRegistry -> VisualBindingSet tetap menjadi jalur efek.",
"TempoSegmentRegistry menyediakan informasi tempo per SongInstance untuk Vinyl BPM Sync. MusicStyleRegistry hanya menghasilkan recipe assignment; tidak mengubah beat math."
]),
("4. BeatPresetRegistry", [
"BeatPresetDefinition memuat preset ID, label, category, description, VisualProperty list, recommended intensity, text_safe, rotation_required, dan ai_aliases.",
"Registry wajib menjadi source of truth yang sama untuk UI, renderer capability, tests, serta action registry AI Agent pada STEP berikutnya."
]),
("5. Katalog 18 Preset", [
"Legacy: subtle_beat_pulse, bass_pulse, strong_punch, onset_flash, rotation_nudge, energy_breathe.",
"Baru: beat_zoom, bass_zoom, glow_pump, bass_glow, strong_glow, beat_tilt, bass_tilt, energy_zoom, energy_glow, club_punch, bass_punch, cinematic_swell.",
"Mapping utama: Beat Zoom BEAT->Zoom +0.035; Bass Zoom BASS->Zoom +0.070; Glow Pump BEAT->Glow +0.220; Bass Glow BASS->Glow +0.420; Strong Glow STRONG->Glow +0.600.",
"Beat Tilt BEAT->Rotation +1.2; Bass Tilt BASS->Rotation +2.5; Energy Zoom ENERGY_UP->Zoom +0.045; Energy Glow ENERGY_UP->Glow +0.350.",
"Club Punch STRONG->Scale +0.110, Zoom +0.090, Glow +0.550; Bass Punch BASS->Scale +0.085, Zoom +0.045, Glow +0.300; Cinematic Swell ENERGY_UP->Scale +0.050, Zoom +0.040, Glow +0.280."
]),
("6. Compatibility", [
"song_cover, vinyl, background, song_visual, spectrum dapat memakai semua 18 preset bila rotation-safe.",
"text/song_title tidak boleh menawarkan rotation_nudge, beat_tilt, atau bass_tilt.",
"Preset rotation-required otomatis difilter bila base rotation bukan 0 atau pivot bukan center."
]),
("7. Music Style", [
"MusicStylePreset: CHILL, AMBIENT, POP, ROCK, EDM, HIP_HOP, DANGDUT_REMIX, ACOUSTIC, CINEMATIC.",
"Chill: cover Energy Breathe 70%; background/song_visual Cinematic Swell 45%; spectrum Glow Pump 40%; vinyl Beat Pulse 45%; title Onset Flash 30%.",
"Ambient: cover Cinematic Swell 55%; background/song_visual Energy Glow 40%; spectrum Energy Breathe 35%; vinyl Beat Pulse 30%; title Energy Glow 30%.",
"Pop: cover Beat Zoom 85%; background/song_visual Beat Pulse 50%; spectrum Glow Pump 75%; vinyl Bass Pulse 70%; title Onset Flash 50%.",
"Rock: cover Bass Punch 100%; background/song_visual Beat Tilt 45%; spectrum Strong Glow 85%; vinyl Strong Punch 80%; title Beat Zoom 65%.",
"EDM: cover Club Punch 115%; background/song_visual Strong Glow 80%; spectrum Bass Punch 110%; vinyl Bass Punch 100% + BPM Sync; title Onset Flash 65%.",
"Hip-Hop: cover Bass Punch 105%; background/song_visual Bass Zoom 55%; spectrum Bass Glow 90%; vinyl Bass Pulse 85% + BPM Sync; title Beat Pulse 45%.",
"Dangdut Remix: cover Club Punch 110%; background/song_visual Beat Zoom 65%; spectrum Bass Punch 105%; vinyl Bass Pulse 100% + BPM Sync; title Onset Flash 55%.",
"Acoustic: cover Beat Pulse 45%; background/song_visual Energy Breathe 30%; spectrum Beat Pulse 30%; vinyl Beat Pulse 30%; title Onset Flash 25%.",
"Cinematic: cover Cinematic Swell 75%; background/song_visual Cinematic Swell 65%; spectrum Energy Glow 45%; vinyl Energy Breathe 35%; title Energy Glow 40%."
]),
("8. Apply Music Style", [
"Apply Music Style hanya menulis layer.animation['beat_v1'] dan optional Vinyl BPM properties. Asset, timing, layer structure, dan audio tidak boleh berubah.",
"Satu Apply dijalankan sebagai satu EditorController transaction sehingga satu Undo mengembalikan seluruh perubahan.",
"Locked layers dilewati dan dilaporkan. Pemilihan dropdown tidak auto-apply."
]),
("9. Vinyl BPM Sync", [
"Vinyl properties: bpm_sync=false, beats_per_rotation=4.0, bpm_sync_min_confidence=0.55; allowed beats_per_rotation 1/2/4/8.",
"Synced spin seconds = 60 * beats_per_rotation / BPM. 120 BPM dan 4 beats/rotation = 2 detik per putaran.",
"Jika BPM=0, quality LOW/SILENT, atau confidence di bawah threshold, fallback ke existing spin_seconds. Tidak boleh memaksa BPM palsu."
]),
("10. TempoSegment Runtime", [
"TempoSegment menyimpan song_id, asset_id, start_tick, end_tick, bpm, confidence, quality.",
"BeatVisualRuntime mendapat tempo_segments dan API tempo_at_tick, vinyl_spin_seconds_at, vinyl_phase_cycles_at.",
"Pada overlap Free Timeline, segment valid dengan start_tick paling baru menjadi tempo visual aktif. Phase sengaja reset per awal lagu."
]),
("11. Runtime Requirement", [
"BPM Sync wajib bekerja walaupun Vinyl tidak memiliki beat_v1. Tambahkan document_needs_beat_runtime = beat assignment ATAU enabled Vinyl bpm_sync.",
"Editor preview, Accurate Preview, dan final render memakai runtime preparation yang sama; analyzer tidak dijalankan di FFmpeg compiler."
]),
("12. Preview & Render", [
"PreviewCanvas Vinyl menghitung deterministic phase dari project tick; tidak memakai wall-clock.",
"Final render menggunakan piecewise expression per TempoSegment. Pada tempo valid phase=(T-song_start)/synced_spin_seconds; fallback=T/spin_seconds.",
"Beat pulse/punch dari STEP09 tetap boleh ditumpuk setelah source Vinyl BPM-sync."
]),
("13. UI", [
"Property Inspector Vinyl mendapat BPM Sync checkbox dan beats-per-rotation selector 1/2/4/8.",
"Editor toolbar mendapat Gaya Beat dropdown dan Terapkan button. Memilih dropdown sendiri tidak memutasi document.",
"Advanced preset combo membaca BeatPresetRegistry, bukan hard-coded label list."
]),
("14. AI-ready Metadata", [
"Registry menyimpan ai_aliases, misalnya club_punch: edm punch/hentak kuat/club beat; bass_punch: bass punch/hentak bass; cinematic_swell: cinematic/naik perlahan/dramatic swell.",
"Music Style juga memiliki aliases; STEP10 hanya menyediakan metadata, bukan eksekusi natural-language."
]),
("15. Safety / Limits", [
"Scale tetap clamp 0.75..1.35, zoom 1.0..1.25, glow 0..1, rotation offset -12..12 derajat.",
"Render cap Beat layers tetap 4; command rows tetap 250000; tempo confidence default minimum 0.55."
]),
("16. Backward Compatibility", [
"Enam preset lama mempertahankan ID dan default mapping.",
"Project beat_v1 STEP09 tetap valid. Project tanpa Beat/BPM Sync tidak menganalisis audio tambahan.",
"Vinyl tanpa bpm_sync mempertahankan spin_seconds persis. Music Style adalah recipe action, bukan schema global wajib."
]),
("17. Test Gate", [
"Focused tests harus mencakup 18 unique preset, legacy mapping unchanged, registry coverage, 9 styles, Undo-safe apply, locked layer skip, timing/assets unchanged.",
"Tempo tests mencakup low-confidence fallback, 120 BPM/4 beats=2s, 90 BPM/4 beats=2.6667s, active segment selection, overlap latest-start, runtime without beat_v1.",
"Preview/render tests mencakup deterministic seek phase, reset per song, fallback spin, piecewise expression, BPM-spin + beat punch coexistence.",
"Regression STEP04-09 wajib PASS. Real FFmpeg target: BPM-spin Vinyl, advanced Cover, advanced Spectrum. Windows UI headless smoke wajib PASS."
]),
("18. Acceptance Gate", [
"G1 STEP10 DOCX committed before source changes; G2 18 presets registry-covered; G3 six legacy mappings unchanged; G4 nine styles deterministic/Undo-safe.",
"G5 UI advanced preset/music style/BPM sync usable; G6 TempoSummary with fail-closed fallback; G7 preview seek-safe; G8 final render BPM spin visibly changes output.",
"G9 BPM spin coexists with beat pulse/punch; G10 AI-ready registry exists; G11 no-beat/no-sync path unchanged; G12 STEP04-09 regression PASS; G13 real FFmpeg targets PASS; G14 stable main unchanged."
]),
("19. Definition of Done", [
"18 Beat Animation presets and 9 Music Style recipes tersedia; registry menjadi metadata tunggal.",
"Vinyl sync BPM per song berdasarkan TempoSummary dan fallback aman pada tempo tidak meyakinkan.",
"UI dapat apply styles dan BPM controls secara Undo-safe; preview/final render deterministic.",
"Advanced effects tidak merusak Spectrum atau audio master; Windows/Ubuntu/real-FFmpeg gates harus lulus."
]),
("20. Handoff STEP 11", [
"STEP11 boleh membangun AI Agent action registry/natural-language mapping ke preset, style, intensity dan BPM sync.",
"Shake/random/particles tetap ditunda karena memerlukan event-phase oscillator atau particle lifecycle, bukan sekadar envelope magnitude STEP06."
]),
]

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def build():
    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.66); sec.right_margin=Inches(.66)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(9)
    for name,size,color in [("Title",23,"17365D"),("Heading 1",14,"17365D"),("Heading 2",11,"2F5597")]:
        doc.styles[name].font.name="Aptos Display"; doc.styles[name].font.size=Pt(size); doc.styles[name].font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 10\nAdvanced Beat Animation Catalog & Music Presets")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Beat Animation Engine V2 — Source of Truth / Detailed Implementation Contract").bold=True
    table=doc.add_table(rows=1, cols=2); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(("Item","Nilai")):
        table.rows[0].cells[i].text=h; shade(table.rows[0].cells[i],"17365D")
        for run in table.rows[0].cells[i].paragraphs[0].runs: run.font.color.rgb=RGBColor(255,255,255); run.bold=True
    for a,b in [
        ("Repository","inoriko920-dev/Full-Album-Maker"),
        ("Development branch","feature/beat-animation-engine-v2"),
        ("Stable main","584c94774e6197ecf60ecedb3bffd5a8797e7737 (v1.5.0)"),
        ("STEP09 verified commit","dd6435cccd2c37bdec7042e3b9c9e7fb0b76aa9e"),
        ("Target","18 presets + 9 Music Styles + Vinyl BPM Sync"),
    ]:
        row=table.add_row().cells; row[0].text=a; row[1].text=b
    doc.add_page_break()
    for title, bullets in SECTIONS:
        doc.add_heading(title, level=1)
        for item in bullets:
            p=doc.add_paragraph(style="List Bullet"); p.add_run(item)
    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 10 • 7 Oct 2026")
        r.font.size=Pt(7); r.font.color.rgb=RGBColor(105,105,105)
    doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    line=f"- `{OUT.name}` — SHA256 `{digest}`"
    lines=[x for x in text.splitlines() if OUT.name not in x]
    idx=next((i for i,x in enumerate(lines) if x.startswith("## Handoff rule")),len(lines))
    lines[idx:idx]=[line,""]
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__":
    build()
