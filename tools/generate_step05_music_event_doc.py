from __future__ import annotations

import hashlib
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / "docs" / "beat-animation"
OUT = DOC_ROOT / "STEP05_MUSIC_EVENT_TIMELINE_DERIVED_EVENT_ENGINE_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT = DOC_ROOT / "SOURCE_OF_TRUTH.md"

BASELINE = "584c94774e6197ecf60ecedb3bffd5a8797e7737"
BRANCH = "feature/beat-animation-engine-v2"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = value
        shade(c, "17365D")
        for p in c.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(8)
    for values in rows:
        row = t.add_row()
        for i, value in enumerate(values):
            row.cells[i].text = str(value)
            for p in row.cells[i].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
    doc.add_paragraph("")


def bullets(doc: Document, values: list[str]) -> None:
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(value)


def code(doc: Document, value: str) -> None:
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, "F7F7F7")
    p = c.paragraphs[0]
    r = p.add_run(value)
    r.font.name = "Consolas"
    r.font.size = Pt(8)
    doc.add_paragraph("")


def callout(doc: Document, title: str, text: str, fill: str = "D9EAF7") -> None:
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.cell(0, 0)
    shade(c, fill)
    p = c.paragraphs[0]
    a = p.add_run(title + " — ")
    a.bold = True
    p.add_run(text)
    doc.add_paragraph("")


def build() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.62)
    sec.bottom_margin = Inches(0.62)
    sec.left_margin = Inches(0.68)
    sec.right_margin = Inches(0.68)

    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.2)
    for name, size, color in [
        ("Title", 24, "17365D"),
        ("Heading 1", 15, "17365D"),
        ("Heading 2", 11.5, "2F5597"),
    ]:
        s = doc.styles[name]
        s.font.name = "Aptos Display"
        s.font.size = Pt(size)
        s.font.color.rgb = RGBColor.from_string(color)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FULL-ALBUM-MAKER")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string("4472C4")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 05\nMusic Event Timeline & Derived Event Engine")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Beat Animation Engine V2 — Implementation Planning / Source of Truth").bold = True

    callout(
        doc,
        "Status STEP",
        "STEP 05 hanya membangun vocabulary event, derived-event engine, dan source→project projection. "
        "Tidak membuat UI animasi, tidak mengubah renderer final, dan tidak melompat ke STEP 06.",
        "E2F0D9",
    )
    table(doc, ["Item", "Nilai"], [
        ("Repository", "inoriko920-dev/Full-Album-Maker"),
        ("Development branch", BRANCH),
        ("Stable baseline main", BASELINE + " (v1.5.0)"),
        ("Target", "Deterministic source-relative Music Event Timeline untuk semua animasi"),
        ("Tanggal", "7 Oktober 2026 (WIB)"),
    ])

    doc.add_heading("1. Tujuan STEP 05", 1)
    doc.add_paragraph(
        "STEP 05 menjembatani AudioAnalysisResult STEP 04 dengan Animation Engine. Analyzer hanya memberi fakta dasar "
        "(beat, onset, energy, bass, mid, high, tempo, quality). Animation layer tidak boleh mendefinisikan sendiri "
        "apa itu beat kuat atau bass hit karena hasil akan berbeda antar efek dan sulit diuji."
    )
    callout(doc, "Prinsip utama", "Satu AudioAnalysisResult → satu Derived Event Engine → satu MusicEventTimeline → dipakai semua animasi.")
    bullets(doc, [
        "Source-relative per audio asset; satu file dianalisis dan diderivasi sekali.",
        "Strength dan confidence selalu 0..1.",
        "Beat-derived events fail-closed pada LOW/SILENT.",
        "Local peak + NMS/debounce mencegah event spam.",
        "Projection ke Packed/Free Timeline tidak mengubah ProjectDocument.",
    ])

    doc.add_heading("2. Scope dan non-scope", 1)
    table(doc, ["Masuk STEP 05", "Ditunda"], [
        ("Raw BEAT / ONSET passthrough", "UI animasi"),
        ("STRONG_BEAT", "Renderer FFmpeg animation"),
        ("BASS_HIT", "Preset genre"),
        ("ENERGY_RISE / ENERGY_FALL", "AI Agent binding"),
        ("Project event projection", "DROP / SECTION_CHANGE / KICK-SNARE classifier"),
    ])

    doc.add_heading("3. Baseline STEP 04", 1)
    table(doc, ["Kontrak", "Nilai"], [
        ("TIMEBASE", "240.000 tick/detik"),
        ("SAMPLE_RATE", "24.000 Hz = 10 tick/sample"),
        ("HOP_LENGTH", "512 = 5.120 tick/frame ≈ 21,333 ms"),
        ("Required curves", "energy, bass, mid, high, onset"),
        ("Quality", "SILENT / LOW / MEDIUM / HIGH"),
        ("Beat safety", "LOW/SILENT sudah tidak menghasilkan raw BEAT"),
    ])

    doc.add_heading("4. Vocabulary event V1", 1)
    table(doc, ["Event", "Sumber", "Quality gate", "Fungsi visual"], [
        ("BEAT", "Raw STEP04", "MEDIUM/HIGH", "Pulse ritmis"),
        ("ONSET", "Raw STEP04", "Non-silent", "Transient/accent"),
        ("STRONG_BEAT", "Beat + onset/energy/bass", "MEDIUM/HIGH", "Punch/zoom/scale"),
        ("BASS_HIT", "Bass local peak + onset", "LOW/MEDIUM/HIGH", "Kick/bass impact"),
        ("ENERGY_RISE", "Energy trend", "Non-silent", "Naik intensitas"),
        ("ENERGY_FALL", "Energy trend", "Non-silent", "Mereda"),
    ])
    callout(
        doc, "BASS_HIT pada LOW quality",
        "Tempo confidence berbeda dengan local bass confidence. Track tanpa BPM stabil tetap dapat memiliki transient bass valid.",
        "FFF2CC",
    )

    doc.add_heading("5. Kontrak data", 1)
    code(doc, '''MUSIC_EVENT_SCHEMA = "full-album-maker-music-event-timeline"
MUSIC_EVENT_SCHEMA_VERSION = 1
DERIVED_EVENT_ENGINE_VERSION = "music-events-v1"

MusicEvent:
  tick: int                 # source-relative
  event_type: BEAT|ONSET|STRONG_BEAT|BASS_HIT|ENERGY_RISE|ENERGY_FALL
  strength: float           # 0..1
  confidence: float         # 0..1
  source: str

MusicEventTimeline:
  asset_id
  content_sha256
  analysis_settings_signature
  analyzer_version
  engine_version
  duration_tick
  quality
  events: tuple[MusicEvent, ...]''')

    doc.add_heading("6. Ordering & determinisme", 1)
    table(doc, ["Aturan", "Keputusan"], [
        ("Sort", "tick ascending"),
        ("Same-tick priority", "STRONG_BEAT → BASS_HIT → BEAT → ONSET → ENERGY_RISE → ENERGY_FALL"),
        ("Tie break", "event type lalu source"),
        ("Float safety", "NaN/Inf ditolak; valid values clamp 0..1"),
        ("Repeatability", "Input + settings sama menghasilkan tuple event sama"),
        ("Mutation", "AudioAnalysisResult immutable dari perspektif engine"),
    ])

    doc.add_heading("7. DerivedEventSettings V1", 1)
    table(doc, ["Parameter", "Default", "Makna"], [
        ("strong_beat_quantile", "0.75", "Adaptive beat threshold"),
        ("strong_beat_min_score", "0.60", "Hard floor"),
        ("weights onset/energy/bass", "0.45 / 0.25 / 0.30", "Beat salience"),
        ("bass_peak_min", "0.58", "Bass hard floor"),
        ("bass_onset_min", "0.20", "Transient concurrence"),
        ("bass_peak_radius_frames", "2", "Local max radius"),
        ("bass_min_separation_ms", "120", "NMS"),
        ("energy_window_ms", "360", "Trend window"),
        ("energy_delta_min", "0.18", "Minimal change"),
        ("energy_level_min", "0.45", "Noise guard"),
        ("energy_event_separation_ms", "800", "Debounce"),
    ])

    doc.add_heading("8. STRONG_BEAT", 1)
    code(doc, '''score = clamp(0.45*onset + 0.25*energy + 0.30*bass)
adaptive = quantile(all_beat_scores, 0.75)
threshold = max(0.60, adaptive)
candidate if score >= threshold
confidence = min(tempo.confidence, 0.55 + 0.45*score)''')
    bullets(doc, [
        "Hanya raw BEAT yang boleh menjadi STRONG_BEAT.",
        "Jika beat < 4, gunakan hard floor tanpa mempercayai quantile.",
        "LOW/SILENT tidak menghasilkan STRONG_BEAT.",
    ])

    doc.add_heading("9. BASS_HIT", 1)
    code(doc, '''candidate frame i:
  bass[i] >= 0.58
  local maximum within ±2 frames
  onset[i] >= 0.20

score = clamp(0.72*bass[i] + 0.28*onset[i])
NMS keeps strongest candidates separated by >=120 ms
confidence = clamp(0.60*bass[i] + 0.40*onset[i])''')
    callout(doc, "Safety", "Bass panjang/plateau tidak boleh menghasilkan hit setiap frame; local peak + onset + NMS wajib.", "FFF2CC")

    doc.add_heading("10. ENERGY_RISE / ENERGY_FALL", 1)
    code(doc, '''window ≈ 360 ms
past = mean(energy[i-window:i])
future = mean(energy[i:i+window])
delta = future - past

RISE: delta >= +0.18 and future >= 0.45
FALL: delta <= -0.18 and past >= 0.45
strength = clamp(abs(delta)/0.45)
same-type debounce >=800 ms, keep strongest''')

    doc.add_heading("11. Raw passthrough", 1)
    table(doc, ["Input", "Output"], [
        ("AnalysisEvent.BEAT", "MusicEventType.BEAT; strength/confidence dipertahankan"),
        ("AnalysisEvent.ONSET", "MusicEventType.ONSET; strength/confidence dipertahankan"),
    ])

    doc.add_heading("12. Source-relative → project projection", 1)
    code(doc, '''include event if source_in_tick <= event.tick < source_out_tick
project_tick = ResolvedSong.start_tick + (event.tick - source_in_tick)''')
    table(doc, ["Kasus", "Perilaku"], [
        ("Packed", "Gunakan ResolvedSong.start_tick"),
        ("Free", "Gunakan posisi hasil TimelineResolver"),
        ("Trim awal/akhir", "Event di luar source range dibuang"),
        ("Asset reused", "Source timeline dipakai ulang untuk setiap SongInstance"),
        ("Crossfade", "Event kedua lagu tetap ada; arbitration ditunda ke STEP06"),
        ("Disabled song", "Tidak diproyeksikan"),
    ])
    code(doc, '''ProjectedMusicEvent:
  project_tick
  source_tick
  song_id
  asset_id
  event_type
  strength
  confidence''')

    doc.add_heading("13. Performance", 1)
    table(doc, ["Operasi", "Target"], [
        ("Strong beat", "O(B)"),
        ("Bass peak scan", "O(F)"),
        ("Energy trend", "O(F) dengan prefix/rolling mean"),
        ("NMS", "O(C log C)"),
        ("Projection", "O(E) per song range"),
    ])
    callout(doc, "Long album", "Dilarang O(F×window) dan dilarang menggandakan seluruh curve menjadi list Python besar.")

    doc.add_heading("14. Cache / memo", 1)
    doc.add_paragraph(
        "Tidak ada cache disk kedua di V1. Derived generation jauh lebih murah daripada decode/FFT/librosa. "
        "Boleh memo in-memory dengan identity content_sha256 + analysis settings + derived settings + engine version."
    )

    doc.add_heading("15. Modul implementasi", 1)
    code(doc, '''music_event_contract.py
  MusicEventType, MusicEvent, MusicEventTimeline, ProjectedMusicEvent, DerivedEventSettings

music_event_engine.py
  build_music_event_timeline()
  derive_strong_beats()
  derive_bass_hits()
  derive_energy_events()

music_event_projection.py
  project_song_events()
  project_album_events()''')
    doc.add_paragraph("Modul harus Qt-independent, FFmpeg-independent, dan librosa-independent.")

    doc.add_heading("16. Validation & failure semantics", 1)
    table(doc, ["Kondisi", "Respons"], [
        ("Invalid AudioAnalysisResult", "Fail fast ValueError"),
        ("Curve length mismatch", "Fail fast"),
        ("NaN/Inf", "Fail fast"),
        ("LOW quality tanpa beats", "Normal; strong beat kosong"),
        ("No bass peaks", "Normal; bass hit kosong"),
        ("SILENT", "Tidak menghasilkan derived events"),
        ("Unknown/inactive song projection", "Empty atau explicit error sesuai API; tidak silent corruption"),
    ])

    doc.add_heading("17. Test matrix", 1)
    tests = [
        "Contract valid/reject negative tick/reject invalid unit values",
        "Deterministic ordering dan same-tick priority",
        "Raw BEAT dan ONSET passthrough",
        "LOW quality suppresses STRONG_BEAT",
        "MEDIUM/HIGH strong beat + adaptive threshold + hard floor",
        "Bass local maximum, plateau anti-spam, onset concurrence, NMS strongest",
        "Energy rise/fall synthetic ramps, flat curve no-event, debounce",
        "SILENT safe empty derived timeline",
        "Same input produces identical event tuple",
        "Packed projection exact tick",
        "source_in/source_out trim exact and end-exclusive",
        "Free timeline projection exact tick",
        "Same asset reused across SongInstances",
        "Disabled song omitted",
        "Crossfade does not shift source-relative tick",
        "Large curve linear-time sanity",
        "Settings change changes identity",
        "AudioAnalysisResult not mutated",
    ]
    for i, value in enumerate(tests, start=1):
        doc.add_paragraph(f"{i:02d}. {value}")

    doc.add_heading("18. CI gate", 1)
    bullets(doc, [
        "Focused STEP05 tests pada ubuntu-latest dan windows-latest.",
        "Regression tests/test_step04_audio_analysis.py tetap PASS.",
        "Tidak merge ke main.",
        "Workflow hanya menyentuh branch feature/beat-animation-engine-v2.",
    ])

    doc.add_heading("19. Urutan implementasi", 1)
    bullets(doc, [
        "Commit DOCX STEP05 + SOURCE_OF_TRUTH terlebih dahulu.",
        "Implement contract + validation.",
        "Implement raw passthrough + stable ordering.",
        "Implement STRONG_BEAT.",
        "Implement BASS_HIT + NMS.",
        "Implement ENERGY_RISE/FALL dengan prefix/rolling mean.",
        "Implement source→project projection.",
        "Tambah focused tests + STEP04 regression.",
        "CI Windows + Ubuntu wajib PASS.",
    ])

    doc.add_heading("20. Event sengaja ditunda", 1)
    table(doc, ["Event", "Alasan"], [
        ("DROP", "Heuristic sederhana terlalu mudah false-positive; perlu konteks pre/post drop lebih kuat."),
        ("SECTION_CHANGE", "Perlu novelty/segmentation yang belum tersedia."),
        ("KICK/SNARE", "Frequency band bukan instrument classifier."),
        ("BAR/DOWNBEAT", "Beat tracker saat ini belum menjamin meter/downbeat."),
    ])
    callout(doc, "Aturan kualitas", "Lebih baik sedikit event yang benar daripada label musik berlebihan yang sebenarnya tidak dideteksi.", "E2F0D9")

    doc.add_heading("21. Definition of Done", 1)
    bullets(doc, [
        "STEP05 DOCX committed dan SOURCE_OF_TRUTH diperbarui.",
        "main tetap baseline v1.5.0.",
        "MusicEventTimeline immutable dan deterministic.",
        "STRONG_BEAT quality-gated.",
        "BASS_HIT local peak + onset + NMS.",
        "ENERGY_RISE/FALL trend + debounce.",
        "Projection trim/Packed/Free exact integer tick.",
        "Tidak ada DROP/SECTION_CHANGE palsu.",
        "Focused STEP05 + STEP04 regression PASS Windows dan Ubuntu.",
        "Tidak ada perubahan UI/render final.",
    ])

    doc.add_heading("22. Acceptance Gate", 1)
    table(doc, ["Gate", "PASS"], [
        ("G1 Source of Truth", "DOCX committed sebelum code"),
        ("G2 Determinism", "Same input/settings → same timeline"),
        ("G3 Safety", "LOW/SILENT tidak memicu strong beat"),
        ("G4 No spam", "NMS/debounce tervalidasi"),
        ("G5 Time mapping", "Trim + Packed + Free exact tick"),
        ("G6 Isolation", "ProjectDocument/STEP04 cache tidak dimutasi"),
        ("G7 Tests", "Focused + STEP04 regression PASS Windows/Ubuntu"),
        ("G8 Stable main", "main masih baseline"),
    ])

    doc.add_heading("23. Handoff ke STEP 06", 1)
    doc.add_paragraph(
        "STEP 06 baru membangun Animation Signal/Envelope Engine: attack, decay, smoothing, intensity scaling, cooldown arbitration, "
        "dan conversion event menjadi parameter frame-by-frame. STEP 06 dilarang mendefinisikan ulang strong beat/bass hit."
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 05 • 7 Oct 2026")
    rr.font.size = Pt(7)
    rr.font.color.rgb = RGBColor(110, 110, 110)

    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    digest = hashlib.sha256(OUT.read_bytes()).hexdigest()
    text = SOT.read_text(encoding="utf-8")
    line = f"- `{OUT.name}` — SHA256 `{digest}`"
    if OUT.name not in text:
        marker = "\n## Handoff rule"
        text = text.replace(marker, "\n" + line + "\n" + marker)
    else:
        lines = []
        for old in text.splitlines():
            if OUT.name in old:
                lines.append(line)
            else:
                lines.append(old)
        text = "\n".join(lines) + ("\n" if text.endswith("\n") else "")
    SOT.write_text(text, encoding="utf-8")
    print(OUT)
    print(digest)


if __name__ == "__main__":
    build()
