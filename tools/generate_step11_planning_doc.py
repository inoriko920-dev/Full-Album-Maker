from __future__ import annotations

import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
DOC_ROOT=ROOT/"docs"/"beat-animation"
OUT=DOC_ROOT/"STEP11_AI_AGENT_ACTION_REGISTRY_NATURAL_LANGUAGE_BEAT_CONTROL_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOC_ROOT/"SOURCE_OF_TRUTH.md"
DARK="17365D"; BLUE="D9EAF7"; GREEN="E2F0D9"; YELLOW="FFF2CC"; RED="FCE4D6"

def shade(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); tcPr.append(shd)

def margins(cell,v=46,h=54):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for n,val in (("top",v),("bottom",v),("start",h),("end",h)):
        node=tcMar.find(qn(f"w:{n}"))
        if node is None:
            node=OxmlElement(f"w:{n}"); tcMar.append(node)
        node.set(qn("w:w"),str(val)); node.set(qn("w:type"),"dxa")

def nosplit(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

def add_table(doc,headers,rows,fill=DARK):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; nosplit(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,fill); margins(c)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(7.75)
    for values in rows:
        row=t.add_row(); nosplit(row)
        for i,value in enumerate(values):
            c=row.cells[i]; c.text=str(value); margins(c)
            for p in c.paragraphs:
                for r in p.runs: r.font.size=Pt(7.45)
    return t

def callout(doc,title,text,fill=BLUE):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"
    c=t.cell(0,0); shade(c,fill); margins(c,72,86)
    p=c.paragraphs[0]; r=p.add_run(title+" — "); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text)

def code(doc,text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; nosplit(t.rows[0])
    c=t.cell(0,0); shade(c,"F7F7F7"); margins(c,56,78)
    r=c.paragraphs[0].add_run(text); r.font.name="Consolas"; r.font.size=Pt(7.15)

def bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(1.6); p.add_run(item)

def nums(doc,items):
    for item in items:
        p=doc.add_paragraph(style="List Number"); p.paragraph_format.space_after=Pt(1.6); p.add_run(item)

def build():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.56); sec.bottom_margin=Inches(.56); sec.left_margin=Inches(.62); sec.right_margin=Inches(.62)
    styles=doc.styles; styles["Normal"].font.name="Aptos"; styles["Normal"].font.size=Pt(8.9)
    for name,size,color in (("Title",22.5,"17365D"),("Heading 1",14.3,"17365D"),("Heading 2",11.0,"2F5597")):
        styles[name].font.name="Aptos Display"; styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor.from_string(color)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.size=Pt(10.8); r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 11\nAI Agent Action Registry & Natural-Language Beat Control")
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Beat Animation Engine V2 — Source of Truth / Detailed Implementation Contract"); r.bold=True; r.font.size=Pt(10.4)
    callout(doc,"Status STEP","STEP 11 menghubungkan AI Agent ke Beat Animation V2 melalui whitelist action, permission beat.write, Preview Diff, dan satu transaction Undo. Provider tidak pernah menulis JSON proyek, filesystem, FFmpeg, atau beat math langsung.",GREEN)
    add_table(doc,["Item","Nilai"],[
        ("Repository","inoriko920-dev/Full-Album-Maker"),
        ("Development branch","feature/beat-animation-engine-v2"),
        ("Stable main","584c94774e6197ecf60ecedb3bffd5a8797e7737 (v1.5.0)"),
        ("STEP 10 verified head","bbe324c6282119017663a3ffd2b263b5ec8b298a"),
        ("Tanggal","7 Oktober 2026 (WIB)"),
        ("AI baseline","STEP09 AgentPlan + Preview Diff + permission + atomic transaction"),
        ("Target","Beat actions aman + canonical local NLU + Gemini function tools"),
    ])
    doc.add_page_break()

    sections=[
      ("1. Tujuan",[
        "AI Agent memahami instruksi Beat Animation Bahasa Indonesia natural.",
        "Mutasi selalu memakai domain command/recipe STEP09-10.",
        "Permission beat.write terpisah dari Visual/Spectrum.",
        "Preview Diff wajib sebelum execute.",
        "Satu plan tetap satu revision dan satu Undo.",
        "Fast-path lokal hanya untuk intent ber-confidence tinggi.",
      ]),
      ("3. Security Boundary",[
        "Hanya stable layer IDs, preset IDs, dan Music Style IDs dari sanitized context.",
        "Tidak boleh mengarang ID/preset/style.",
        "Tidak ada shell/filesystem/arbitrary code/render/analyzer action.",
      ]),
    ]
    for title,items in sections:
        doc.add_heading(title,level=1); bullets(doc,items)

    doc.add_heading("2. Contoh UX Target",level=1)
    add_table(doc,["Perintah pengguna","Rencana aman"],[
        ("Buat animasinya cocok untuk EDM","apply_music_style(style_id='edm')"),
        ("Pakai Dangdut Remix","apply_music_style(style_id='dangdut_remix')"),
        ("Pakai Club Punch 120% di layer ini","set_beat_preset(selected layer, club_punch, 1.20)"),
        ("Bass-nya lebih kuat","adjust_beat_intensity(selected beat layer, +0.25)"),
        ("Matikan beat di layer ini","clear_beat_animation(selected layer)"),
        ("Sinkronkan vinyl ke BPM","set_vinyl_bpm_sync(selected/unique vinyl, true, 4 beats)"),
    ])
    callout(doc,"Ambiguity rule","Perintah layer-spesifik tanpa target unik harus meminta klarifikasi; agent tidak boleh memilih layer acak.",YELLOW)

    doc.add_heading("4. Permission Baru",level=1)
    code(doc,'AgentPermission.BEAT_WRITE = "beat.write"\nAI Context Dock menambah checkbox: Beat Animation')
    bullets(doc,["Semua action Beat STEP11 membutuhkan beat.write.","Grant divalidasi sebelum dry-run mutation.","Plan multi-domain tetap meminta union permission."])

    doc.add_heading("5. Lima Action Contract",level=1)
    add_table(doc,["Action","Argumen","Efek"],[
        ("set_beat_preset","layer_ids, preset_id, intensity","Set satu preset Beat"),
        ("adjust_beat_intensity","layer_ids, delta","Ubah intensity relatif tanpa ganti preset"),
        ("clear_beat_animation","layer_ids","Hapus beat_v1"),
        ("apply_music_style","style_id","Terapkan recipe STEP10"),
        ("set_vinyl_bpm_sync","layer_ids, enabled, beats_per_rotation","Atur BPM Sync Vinyl"),
    ])

    doc.add_heading("6. set_beat_preset",level=1)
    code(doc,"Validation:\n- layer_id ada di writable Beat Context\n- layer tidak locked\n- preset ada di registry\n- preset compatible\n- intensity finite 0..2\n\nCommand: SetLayerAnimationValue(layer_id, 'beat_v1', payload)")
    doc.add_heading("7. adjust_beat_intensity",level=1)
    code(doc,"new_intensity = clamp(current + delta, 0, 2)\ndelta allowed -1..+1\nexisting preset list dipertahankan")
    bullets(doc,["'sedikit lebih kuat' +0.10; 'lebih kuat' +0.25; 'jauh lebih kuat' +0.50.","Frasa lembut/pelan memakai delta negatif."])

    doc.add_heading("8. apply_music_style",level=1)
    doc.add_paragraph("Resolver memakai music_style_presets.apply_music_style() dan ReplaceDocument sehingga recipe kompleks tetap satu domain command dan satu Undo.")
    doc.add_heading("9. set_vinyl_bpm_sync",level=1)
    code(doc,"layer_ids: stable vinyl IDs\nenabled: bool\nbeats_per_rotation: 1 | 2 | 4 | 8\n\nSetLayerProperty('bpm_sync', enabled)\nSetLayerProperty('beats_per_rotation', value)")
    doc.add_heading("10. clear_beat_animation",level=1)
    doc.add_paragraph("Menghapus hanya beat_v1 melalui SetLayerAnimationValue(..., _missing=True); BPM Sync tidak ikut dimatikan kecuali diminta.")

    doc.add_heading("11. Beat Context Sanitized",level=1)
    code(doc,'payload["beat_context"] = {\n  "preset_catalog": [{id,label,category}, ...],\n  "music_styles": [{id,label}, ...],\n  "layers": [{layer_id,type,selected,locked,current_presets,intensity,supported_presets,bpm_sync,beats_per_rotation}],\n  "writable_layer_ids": [...]\n}')
    callout(doc,"Privacy","Tidak membawa locator/path media, API key, raw audio curve, cache path, atau FFmpeg command file.",GREEN)

    doc.add_heading("12. Writable Layer Scope",level=1)
    bullets(doc,[
        "Jika ada selected Beat-capable layers, target default hanya selection.",
        "Tanpa selection, layer-specific intent harus klarifikasi kecuali target compatible hanya satu.",
        "Project-wide apply_music_style tidak memerlukan layer_ids.",
        "Prompt eksplisit semua layer dapat memakai IDs dari context dan tetap divalidasi resolver.",
    ])

    doc.add_heading("13. Deterministic Local NLU Fast-Path",level=1)
    add_table(doc,["Pattern jelas","Action"],[
        ("terapkan/pakai + exact Music Style alias","apply_music_style"),
        ("sinkronkan vinyl ke bpm/tempo","set_vinyl_bpm_sync"),
        ("matikan beat/animasi beat","clear_beat_animation"),
        ("exact preset alias + optional N%","set_beat_preset"),
        ("lebih kuat/lembut pada selected Beat layer","adjust_beat_intensity"),
    ])
    bullets(doc,["Token/alias-based; tidak fuzzy-match luas.","Ambiguity menghasilkan clarification.","Fast-path tetap membuat AgentPlan normal sehingga permission/Preview Diff wajib."])

    doc.add_heading("14. Gemini Function Tool Extension",level=1)
    bullets(doc,[
        "Tambahkan lima function declaration Beat.",
        "Gunakan apply_music_style untuk gaya genre/proyek.",
        "Gunakan set_beat_preset untuk preset eksplisit.",
        "Gunakan adjust_beat_intensity untuk lebih kuat/lembut tanpa mengganti preset.",
        "Gunakan set_vinyl_bpm_sync untuk Vinyl.",
        "Jangan mengarang preset/style; gunakan Beat Context.",
    ])
    doc.add_heading("15. Natural Language Normalization",level=1)
    add_table(doc,["Bahasa pengguna","Canonical ID"],[
        ("EDM, club, dance","edm"),("dangdut remix, koplo remix","dangdut_remix"),("hip hop, hiphop, rap beat","hip_hop"),
        ("cinematic, dramatis","cinematic"),("hentak bass, bass punch","bass_punch"),("hentak kuat, EDM punch","club_punch"),("naik perlahan, dramatic swell","cinematic_swell"),
    ])

    doc.add_heading("16. Preview Diff & Impact",level=1)
    bullets(doc,["Beat assignment tercatat sebagai changed_layer_ids.","Music Style dapat mengubah banyak layer tetapi recipe tetap satu action-domain command.","BPM Sync hanya memengaruhi properties target.","Plan card harus human-readable."])
    doc.add_heading("17. AI Workspace Plan Card",level=1)
    add_table(doc,["Action","Plan card"],[
        ("apply_music_style","Gaya Beat — EDM / Club"),("set_beat_preset","Beat Preset — Club Punch • 120%"),
        ("adjust_beat_intensity","Beat Intensity — +25%"),("clear_beat_animation","Beat Animation — nonaktifkan"),
        ("set_vinyl_bpm_sync","Vinyl BPM Sync — aktif • 4 beat/putaran"),
    ])

    doc.add_heading("18. Failure Semantics",level=1)
    add_table(doc,["Kondisi","Respons"],[
        ("Unknown preset/style","Reject before mutation"),("Layer locked","Reject"),("Preset incompatible","Reject + reason"),
        ("No unique layer target","Clarification"),("beat.write off","Permission error"),("Stale revision/fingerprint","Interpret ulang"),
        ("Duplicate plan_id","No second mutation"),("Intensity overflow","Clamp 0..2"),
    ])
    doc.add_heading("19. Undo / Atomicity",level=1)
    doc.add_paragraph("Semua action dalam satu response dikumpulkan Step09TransactionEngine dan dikomit satu revision. Undo AI mengembalikan keadaan sebelum seluruh plan bila belum ada manual edit sesudahnya.")
    doc.add_heading("20. No-Direct-Analyzer Rule",level=1)
    callout(doc,"AI boundary","Agent hanya mengubah konfigurasi Beat/Style/BPM Sync. Tidak memanggil AudioAnalysisService, membaca curve/BPM mentah, atau menentukan timestamp beat.",RED)

    doc.add_heading("21. File Plan",level=1)
    code(doc,"modify:\n  ai_agent_core_step09.py\n  ai_action_registry_step09.py\n  ai_provider_step09.py\n  ai_workspace_step09.py\n  ai_feature_step09.py\n\nnew:\n  ai_beat_nlu_step11.py\n\ntests:\n  test_step11_ai_beat_actions.py\n  test_step11_ai_beat_nlu.py\n  test_step11_ai_beat_provider.py")
    doc.add_heading("22. Test Matrix",level=1)
    tests=[
        "beat.write exists","Beat permission checkbox exists","Beat Context contains 18 presets","Beat Context contains 9 styles",
        "Beat Context contains no locator/path/API key","Selected Beat layer becomes writable","Ambiguous no-selection target clarifies",
        "set_beat_preset valid","incompatible rejected","locked rejected","adjust intensity +0.25","multi-preset preserved",
        "adjust clamps 0..2","adjust without assignment rejected","clear removes only beat_v1","clear leaves BPM Sync unchanged",
        "apply EDM style via ReplaceDocument","apply Dangdut Remix","Music Style one command domain","BPM sync valid vinyl",
        "BPM sync non-vinyl rejected","invalid beats_per_rotation rejected","Beat action requires beat.write","plan denied without permission",
        "dry-run non-destructive","execute one revision","one Undo restores project","duplicate plan no second mutation",
        "stale Beat plan rejected","EDM maps locally","Dangdut Remix maps locally","sync vinyl maps locally","matikan beat maps locally",
        "Club Punch 120 maps locally","lebih kuat maps +0.25","ambiguous phrase clarifies/falls through",
        "Gemini has five Beat tools","Gemini forbids invented preset/style","provider rejects unregistered function","plan card Beat actions",
        "STEP10 regression","STEP09 Agent regression","STEP09 Beat regression","STEP08 regression","STEP07 regression","STEP06 regression","STEP05 regression","STEP04 regression",
    ]
    add_table(doc,["ID","Test"],[(f"{i:02d}",x) for i,x in enumerate(tests,1)])

    doc.add_heading("23. CI Gate",level=1)
    bullets(doc,["Ubuntu + Windows focused STEP11 tests.","STEP09 Agent Core regression PASS.","Beat Engine STEP04-10 regression PASS.","Canonical NLU tests tanpa network.","Gemini diuji via payload/function schema mock."])
    doc.add_heading("24. Implementation Order",level=1)
    nums(doc,["Commit STEP11 DOCX + SOURCE_OF_TRUTH.","Tambah beat.write + Beat Context.","Tambah 5 action resolver.","Tambah deterministic Beat NLU.","Extend Gemini tools/system.","Extend workspace permission/plan cards.","Focused tests.","Regression Windows/Ubuntu.","Audit main unchanged."])

    doc.add_heading("25. Acceptance Gate",level=1)
    add_table(doc,["Gate","PASS Criteria"],[
        ("G1 Source of Truth","STEP11 DOCX committed before code"),("G2 Permission","beat.write enforced"),("G3 Registry","5 Beat actions fail-closed"),
        ("G4 Context","bounded/path-free"),("G5 NLU","canonical deterministic + ambiguity clarifies"),("G6 Gemini","Beat tools + whitelist"),
        ("G7 Preview Diff","accurate impact"),("G8 Atomicity","one revision / one Undo"),("G9 Safety","no analyzer/render/filesystem"),
        ("G10 UI","Beat permission + cards"),("G11 Regression","Agent + STEP04-10 PASS Windows/Ubuntu"),("G12 Stable Main","v1.5.0 baseline"),
    ])
    callout(doc,"Gate keputusan","Semua G1-G12 wajib PASS sebelum STEP11 ditutup.",GREEN)
    doc.add_heading("26. Definition of Done",level=1)
    bullets(doc,["Natural Beat commands bekerja untuk intent dasar.","18 preset + 9 styles menjadi vocabulary registry-backed.","Intensity dapat diubah tanpa mengganti preset.","Vinyl BPM Sync dapat diatur AI.","Semua mutasi lewat action registry + Preview Diff + permission + Undo.","Kasus ambigu tidak mengedit.","No beat math/audio analysis access.","Windows/Ubuntu hijau."])
    doc.add_heading("27. Handoff ke STEP 12",level=1)
    doc.add_paragraph("STEP12 dapat menambah event-phase motion seperti alternating wobble, deterministic camera shake, spark/particle burst, dan directional patterns tanpa menulis ulang boundary AI STEP11.")

    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 11 • 7 Oct 2026"); r.font.size=Pt(7); r.font.color.rgb=RGBColor(105,105,105)

    DOC_ROOT.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    entry=f"- `{OUT.name}` — SHA256 `{digest}`"
    lines=[line for line in text.splitlines() if OUT.name not in line]
    insert_at=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(insert_at,entry); lines.insert(insert_at+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__":
    build()
