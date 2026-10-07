from __future__ import annotations

import base64
import hashlib
import json
import zlib
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1] if Path(__file__).resolve().parent.name == "tools" else Path.cwd()
DOC_ROOT = ROOT / "docs" / "beat-animation"
PAYLOAD_ROOT = DOC_ROOT / "_payload"
PARTS = [PAYLOAD_ROOT / f"planning_docs_payload.part{i}.txt" for i in range(1, 5)]
EXPECTED_PART_SHA256 = {
    1: "932b69b8d88b89361f0001b8bfd81940115b548c940e633162cd370949f1d743",
    2: "b0fb8ad98dc91424c280fab73a81167dc453a94450ad7e290762454231c9474c",
    3: "1ad2c7b6c167e3fb16922b7e1bbf7343162c81cdcdd3d782b20a9ecb3906336f",
    4: "f47d57b1aca63bc432a39e047be84b23ff18a1f68a1e84262ca941951aa055ae",
}
BASELINE_SHA = "584c94774e6197ecf60ecedb3bffd5a8797e7737"
BRANCH = "feature/beat-animation-engine-v2"


def verify_and_load() -> dict[str, list[dict]]:
    encoded = []
    for idx, path in enumerate(PARTS, start=1):
        raw = path.read_bytes()
        actual = hashlib.sha256(raw).hexdigest()
        expected = EXPECTED_PART_SHA256[idx]
        if actual != expected:
            raise RuntimeError(f"Payload part {idx} checksum mismatch: {actual} != {expected}")
        encoded.append(raw.decode("ascii"))
    packed = base64.b64decode("".join(encoded), validate=True)
    decoded = zlib.decompress(packed)
    data = json.loads(decoded.decode("utf-8"))
    if not isinstance(data, dict) or len(data) != 6:
        raise RuntimeError("Planning payload must contain exactly six documents.")
    return data


def shade_cell(cell, fill: str = "D9EAF7") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    heading_specs = {
        "Title": (22, True, "17365D"),
        "Heading 1": (16, True, "17365D"),
        "Heading 2": (12.5, True, "2F5597"),
        "Heading 3": (10.5, True, "365F91"),
    }
    for name, (size, bold, color) in heading_specs.items():
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name in {"Title", "Heading 1"} else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
    if "CodeBlock" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("CodeBlock", 1)
        code.font.name = "Cascadia Mono"
        code.font.size = Pt(8)


def choose_style(doc: Document, raw: str) -> str:
    names = {s.name for s in doc.styles}
    if raw in names:
        return raw
    normalized = raw.casefold().strip()
    if normalized.startswith("heading 1"):
        return "Heading 1"
    if normalized.startswith("heading 2"):
        return "Heading 2"
    if normalized.startswith("heading 3"):
        return "Heading 3"
    if "code" in normalized:
        return "CodeBlock"
    if "title" in normalized:
        return "Title"
    return "Normal"


def build_doc(filename: str, blocks: list[dict]) -> Path:
    doc = Document()
    configure_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7)
    sec.right_margin = Inches(0.7)

    first_nonblank = True
    for block in blocks:
        kind = block.get("t")
        if kind == "p":
            text = str(block.get("text") or "")
            style = choose_style(doc, str(block.get("style") or "Normal"))
            p = doc.add_paragraph(style=style)
            p.add_run(text)
            if first_nonblank and text.strip():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_nonblank = False
            if style == "CodeBlock":
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
        elif kind == "table":
            rows = block.get("rows") or []
            if not rows:
                continue
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r_idx, raw_row in enumerate(rows):
                for c_idx in range(width):
                    value = raw_row[c_idx] if c_idx < len(raw_row) else ""
                    cell = table.cell(r_idx, c_idx)
                    cell.text = str(value or "")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.name = "Aptos"
                            run.font.size = Pt(8.2)
                            if r_idx == 0:
                                run.bold = True
                    if r_idx == 0:
                        shade_cell(cell)
                if r_idx == 0:
                    set_repeat_table_header(table.rows[0])
            doc.add_paragraph("")

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"Full-Album-Maker • Beat Animation Engine V2 • baseline {BASELINE_SHA[:12]}")
    for run in footer.runs:
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor(100, 100, 100)

    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    out = DOC_ROOT / filename
    doc.save(out)
    return out


def main() -> None:
    payload = verify_and_load()
    generated = [build_doc(filename, blocks) for filename, blocks in payload.items()]
    lines = [
        "# Beat Animation Engine V2 — Source of Truth",
        "",
        f"- Baseline stable SHA: `{BASELINE_SHA}`",
        f"- Development branch: `{BRANCH}`",
        "- Stable `main` is not modified by this planning-generation workflow.",
        "- These DOCX files are reconstructed deterministically from checksummed payload parts.",
        "",
        "## Planning documents",
    ]
    for path in generated:
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        lines.append(f"- `{path.name}` — SHA256 `{digest}`")
    lines += [
        "",
        "## Handoff rule",
        "Before changing Beat Animation Engine code, read the master plan and STEP documents in order. Do not skip gates, and keep production `main` unchanged until an explicit merge/release decision.",
        "",
    ]
    (DOC_ROOT / "SOURCE_OF_TRUTH.md").write_text("\n".join(lines), encoding="utf-8")
    print(f"Generated {len(generated)} planning DOCX files.")


if __name__ == "__main__":
    main()
