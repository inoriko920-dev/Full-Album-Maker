from __future__ import annotations

import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parents[1]
DOC_ROOT=ROOT/"docs"/"beat-animation"
OUT=DOC_ROOT/"STEP15_WINDOWS_PORTABLE_FINAL_RELEASE_CANDIDATE_GATE_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOC_ROOT/"SOURCE_OF_TRUTH.md"
DARK="17365D"; GREEN="E2F0D9"; YELLOW="FFF2CC"; RED="FCE4D6"

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),fill); pr.append(shd)
def table(doc,heads,rows):
    t=doc.add_table(rows=1,cols=len(heads)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; c.text=h; shade(c,DARK)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for vals in rows:
        row=t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i,v in enumerate(vals):
            row.cells[i].text=str(v)
            for p in row.cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(8)
def callout(doc,title,text,fill=GREEN):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,fill)
    p=c.paragraphs[0]; r=p.add_run(title+" — "); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text)
def bullet(doc,text):
    p=doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after=Pt(2); p.add_run(text)
def code(doc,text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; c=t.cell(0,0); shade(c,"F7F7F7")
    r=c.paragraphs[0].add_run(text); r.font.name="Consolas"; r.font.size=Pt(7.5)

def build():
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(9)
    for name,size in [("Title",23),("Heading 1",15),("Heading 2",11.5)]:
        doc.styles[name].font.name="Aptos Display"; doc.styles[name].font.size=Pt(size); doc.styles[name].font.color.rgb=RGBColor.from_string(DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 15\nWindows Portable & Final Release Candidate Gate")
    callout(doc,"Status","Final Software Factory gate. No new features; build Windows RC, isolated portable smoke, Beat frozen runtime smoke, real Beat export, integrity evidence, and GO/NO-GO.")
    table(doc,["Item","Value"],[
        ("Development branch","feature/beat-animation-engine-v2"),
        ("Stable main","584c94774e6197ecf60ecedb3bffd5a8797e7737 — v1.5.0"),
        ("STEP14 baseline","b7046c0a2e2b66a05a9607e8df4bbeae61aafd03"),
        ("Release candidate","v1.6.0 RC1"),
        ("Artifact","Full-Album-Maker-v1.6.0-RC1-Windows-Portable.zip"),
        ("Policy","Artifact only; no merge/tag/public release without explicit user instruction"),
    ])
    doc.add_page_break()
    doc.add_heading("1. Version & Freeze",1)
    doc.add_paragraph("v1.6.0 is a backward-compatible minor release over v1.5.0. Beat Animation V2 adds major capability without intentional ProjectDocument breaking migration.")
    for x in ["__version__ and pyproject must both be 1.6.0.","RELEASE_NOTES_v1.6.0.md must exist.","No new visual/AI feature after RC; only evidence-backed bug fixes.","Every bug fix requires a fresh full STEP15 run."]: bullet(doc,x)
    doc.add_heading("2. Clean Windows RC Build",1)
    code(doc,"checkout exact feature head\nPython 3.12.10 / pip 26.2.1\ninstall build/requirements-windows.lock\npip check\npinned FFmpeg+SHA256 / pinned Noto Sans\nfull pytest\nPyInstaller --clean --windowed --onedir")
    callout(doc,"Isolation","RC must come from a clean GitHub Windows runner, never a local developer environment.",YELLOW)
    doc.add_heading("3. RC Workflow Outputs",1)
    table(doc,["Output","Purpose"],[
        ("Full-Album-Maker-v1.6.0-RC1-Windows-Portable.zip","User-testable RC"),
        ("SHA256SUMS-RC1.txt","ZIP integrity"),
        ("RC_MANIFEST.json","Source SHA/version/build provenance"),
    ])
    doc.add_heading("4. Full Regression",1)
    for x in ["python -m pytest -q PASS","python -m compileall -q src PASS","pip check PASS","STEP04-14 regression PASS","Version/GUI title = v1.6.0"]: bullet(doc,x)
    doc.add_heading("5. Extracted ZIP Isolation Smoke",1)
    code(doc,"Extract to Unicode + spaces + apostrophe path\nunset GEMINI_API_KEY / GOOGLE_API_KEY / PYTHONHOME / PYTHONPATH\nPATH=System32 only\nassert no global python\nassert no global ffmpeg")
    for x in ["Run Full Album Maker.exe --portable-smoke.","Verify real short A/V render and production FoundationMainWindow.","portable-smoke.json ok=true, API key absent, audio+video streams present."]: bullet(doc,x)
    doc.add_heading("6. Frozen Beat Runtime Smoke",1)
    for x in ["Run Full Album Maker.exe --beat-runtime-smoke from extracted ZIP.","Frozen numpy/scipy/librosa/numba versions must match pinned runtime.","Synthetic 120 BPM analyze_pcm octave error <5 BPM.","No global Python."]: bullet(doc,x)
    doc.add_heading("7. Real Beat Export RC Smoke",1)
    table(doc,["Check","PASS"],[
        ("Analysis","Bundled runtime analyzes synthetic click audio"),
        ("Cover","Beat preset visibly active"),
        ("Vinyl","BPM Sync active"),
        ("Spectrum","Audio-reactive + Beat modulation"),
        ("Output","MP4 non-empty with audio+video"),
    ])
    doc.add_heading("8. Artifact Integrity",1)
    code(doc,"SHA256SUMS-RC1.txt -> ZIP digest\nRC_MANIFEST.json -> release_candidate/source_commit/app_version/zip_name/zip_sha256/zip_bytes/python/ffmpeg/smokes")
    bullet(doc,"Manifest source_commit must equal GITHUB_SHA.")
    bullet(doc,"Checksum is recomputed after ZIP creation.")
    doc.add_heading("9. Release Notes v1.6.0",1)
    table(doc,["Area","Summary"],[
        ("Beat Analysis","Real beat/onset/energy/bass-mid-high analysis + cache"),
        ("Visual","18 Beat presets + 9 Music Styles"),
        ("Motion","6 motion effects + Spark Burst"),
        ("Combos","8 visual+motion combinations"),
        ("Vinyl","Confidence-gated BPM Sync"),
        ("AI Agent","Beat/motion/style controls with permission, Preview Diff, Undo"),
        ("Hardening","Cancellation, corrupt-cache recovery, runtime memo, batching, long-project guards"),
    ])
    doc.add_heading("10. Merge / Release Policy",1)
    callout(doc,"No automatic merge","Even a PASS result does not move the feature branch to main or publish a GitHub Release. Publication requires explicit user instruction.",RED)
    doc.add_heading("11. Acceptance Gate",1)
    table(doc,["Gate","PASS criteria"],[
        ("G1 Source of Truth","STEP15 DOCX committed before release-gate edits"),
        ("G2 Version","1.6.0 consistent"),
        ("G3 Full Tests","All tests PASS"),
        ("G4 Clean Build","Windows PyInstaller clean runner PASS"),
        ("G5 Isolation","Extracted ZIP works without global Python/FFmpeg/API"),
        ("G6 Frozen Beat","Scientific import + analyze_pcm PASS"),
        ("G7 Real Beat Export","Beat V2 portable export PASS"),
        ("G8 Integrity","SHA256 + manifest + source SHA consistent"),
        ("G9 Artifact","RC ZIP uploaded/downloadable"),
        ("G10 Stable Main","main remains v1.5.0"),
    ])
    callout(doc,"Final decision","G1-G10 PASS = GO for merge/release; any failure = NO-GO and STEP15 remains open.")
    doc.add_heading("12. Definition of Done",1)
    for x in ["Windows RC is built from exact clean branch head.","Beat V2 scientific runtime is proven inside frozen EXE.","Portable manual/Beat workflow works offline.","RC has checksum + provenance manifest.","Release notes are ready.","GO/NO-GO is recorded."]: bullet(doc,x)
    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP15 • RC1 • 7 Oct 2026"); r.font.size=Pt(7)
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    entry=f"- `{OUT.name}` — SHA256 `{digest}`"
    lines=[ln for ln in text.splitlines() if OUT.name not in ln]
    idx=next((i for i,ln in enumerate(lines) if ln.startswith("## Handoff rule")),len(lines))
    lines.insert(idx,entry); lines.insert(idx+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__": build()
