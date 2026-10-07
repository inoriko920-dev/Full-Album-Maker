from __future__ import annotations

import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / "docs" / "beat-animation"
OUT = DOC_ROOT / "STEP08_PREVIEW_RENDER_INTEGRATION_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT = DOC_ROOT / "SOURCE_OF_TRUTH.md"

SECTIONS = [
("1. Tujuan STEP 08", [
"Hubungkan Beat Animation Engine V2 ke preview cepat, Preview Akurat, dan final render tanpa menduplikasi matematika beat.",
"Project tanpa assignment Beat V2 harus tetap identik dengan perilaku lama.",
"Analyzer tidak boleh dijalankan secara tersembunyi di FFmpeg compiler."
]),
("2. Arsitektur", [
"ProjectDocument hanya menyimpan assignment preset pada layer.animation['beat_v1'].",
"BeatRuntimePreparer memastikan cache analysis di luar compiler, lalu membangun MusicEventTimeline, ProjectedMusicEvent, AnimationSignalProgram dan VisualBindingEngine.",
"Fast Preview, Accurate Preview, dan final render membaca runtime yang sama."
]),
("3. Persistence Assignment V1", [
"Format: {'enabled': true, 'presets': ['subtle_beat_pulse','bass_pulse','strong_punch']}.",
"Tidak ada schema bump; runtime/cache/signal program tidak disimpan ke ProjectDocument.",
"Preset tidak dikenal fail closed dengan pesan yang jelas."
]),
("4. Runtime Preparation", [
"Scan beat-enabled layers, collect audio aktif, ensure cache melalui AudioAnalysisService di orchestration layer, build event timeline per asset, project ke album, build signal program, build binding engines per layer.",
"Fast preview tidak boleh memblok UI; saat runtime belum siap ia neutral."
]),
("5. Fast Preview", [
"PreviewCanvas menerima optional BeatVisualRuntime.",
"Base Transform tetap data authoring; effective state hanya dipakai saat paint.",
"Selection/gesture tetap mengedit base transform.",
"Glow preview cepat hanya approximation; Preview Akurat menjadi pixel authority."
]),
("6. Accurate Preview", [
"Evaluate exact tick, bake effective state ke clone ProjectDocument, nonaktifkan beat assignment pada clone agar tidak double-apply, lalu compile frame dengan compiler render yang sama.",
"Original ProjectDocument tidak boleh berubah."
]),
("7. Final Render Strategy", [
"Gunakan FFmpeg sendcmd, bukan expression raksasa berisi ribuan beat.",
"Runtime disampling pada active signal windows; command mengubah scale width/height, rotate angle, dan eq brightness/saturation.",
"Overlay position memakai expression berbasis overlay_w/overlay_h sehingga pivot tetap stabil ketika ukuran berubah."
]),
("8. Render Scope V1", [
"song_cover: full core V1.",
"vinyl: full core V1.",
"background: scale/zoom/glow core dengan rotation restrictions.",
"song_visual/spectrum/text: runtime/preview terlebih dahulu; adapter final render ditunda."
]),
("9. Control Sampling", [
"Control rate = min(project fps, 30 Hz).",
"Sample hanya union active envelope windows, emit neutral reset setelah window, skip perubahan di bawah epsilon.",
"Hard cap command rows 250000 dan maksimal 4 beat-enabled render layers."
]),
("10. Pivot Safety", [
"Overlay x = pivot_project_x*main_w - pivot_x*overlay_w + x_offset*main_w.",
"Overlay y = pivot_project_y*main_h - pivot_y*overlay_h + y_offset*main_h.",
"Dynamic rotation V1 final render dijamin untuk center pivot; unsupported case fail closed."
]),
("11. FFmpeg Runtime Control", [
"scale mendukung runtime width/height command.",
"rotate mendukung runtime angle command.",
"overlay x/y dievaluasi per frame dan dapat dikontrol.",
"eq mendukung brightness/saturation command.",
"Filter instance diberi filter@id agar sendcmd menarget instance yang benar."
]),
("12. Analysis Failure Semantics", [
"No assignment => skip Beat runtime.",
"Cache miss => analysis worker sebelum compile.",
"Analysis failure/source changed => preview/render error actionable.",
"LOW/SILENT tetap mengikuti fail-closed STEP04/05."
]),
("13. Backward Compatibility", [
"No beat assignment tidak boleh mengubah graph lama.",
"Existing Spectrum audio-reactive path tidak diubah.",
"No mandatory dependency baru di luar STEP04 stack.",
"Stable main tidak disentuh."
]),
("14. Module Plan", [
"beat_animation_assignment.py; beat_visual_runtime.py; beat_render_control.py.",
"Modify preview_scene.py, preview_service.py, render_service_v2.py, render_graph.py / Step08FFmpegCompiler hook.",
"Tests: runtime, preview, render control, real FFmpeg."
]),
("15. Tests", [
"Assignment parser, runtime cache preparation, no-mutation, preview exact/approx, snapshot no-double-apply, deterministic control sampling, pivot expressions, command cap, graph backward compatibility, cover/vinyl injection, regression STEP04-07, real FFmpeg sendcmd scale/rotate/eq."
]),
("16. Acceptance Gate", [
"Source-of-truth committed before code.",
"Fast preview core geometry works non-destructively.",
"Accurate preview exact snapshot.",
"Final render song_cover/vinyl physically changes with Beat V2.",
"No-assignment graph unchanged.",
"Regression STEP04-07 PASS Windows/Ubuntu and real FFmpeg smoke PASS.",
"main remains baseline."
]),
("17. Handoff ke STEP 09", [
"STEP09 dapat memperluas adapter ke song_visual/background/text/spectrum modulation dan menambah UI controls. Beat math tetap milik STEP04-07."
]),
]

def build():
    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    doc=Document()
    sec=doc.sections[0]
    sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
    sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(9.5)
    for name,size,color in [("Title",22,"17365D"),("Heading 1",15,"17365D")]:
        st=doc.styles[name]; st.font.name="Aptos Display"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 08\nPreview & Render Integration")
    doc.add_paragraph("Branch: feature/beat-animation-engine-v2\nStable main: 584c94774e6197ecf60ecedb3bffd5a8797e7737\nTanggal: 7 Oktober 2026 (WIB)")
    doc.add_heading("Core Contract", level=1)
    doc.add_paragraph("Audio Analysis → Music Events → Animation Signal → Visual Binding → Preview/Render. Preview dan renderer dilarang mendefinisikan beat/envelope sendiri.")
    for title, items in SECTIONS:
        doc.add_heading(title, level=1)
        for item in items:
            doc.add_paragraph(item, style="List Bullet")
    footer=doc.sections[0].footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 08 • 7 Oct 2026")
    doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    lines=[line for line in text.splitlines() if OUT.name not in line]
    entry=f"- `{OUT.name}` — SHA256 `{digest}`"
    idx=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(idx,entry); lines.insert(idx+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__=="__main__":
    build()
