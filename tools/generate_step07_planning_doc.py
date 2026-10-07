from __future__ import annotations

import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / 'docs' / 'beat-animation'
OUT = DOC_ROOT / "STEP07_VISUAL_PROPERTY_BINDING_CORE_BEAT_ANIMATIONS_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT = DOC_ROOT / 'SOURCE_OF_TRUTH.md'

DARK="17365D"; BLUE="D9EAF7"; GREEN="E2F0D9"; YELLOW="FFF2CC"; RED="FCE4D6"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def margins(cell, v=55, h=65):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar=OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for name,val in [('top',v),('bottom',v),('start',h),('end',h)]:
        n=tcMar.find(qn(`w:${name}`))
        if n is None: n=OxmlElement(`w:${name}`); tcMar.append(n)
        n.set(qn('w:w'),str(val)); n.set(qn('w:type'),'dxa')

def nosplit(row):
    p=row._tr.get_or_add_trPr(); p.append(OxmlElement('w:cantSplit'))

def table(doc, headers, rows, fill=DARK):
    t=doc.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    nosplit(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h; shade(c,fill); margins(c)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for rowvals in rows:
        row=t.add_row(); nosplit(row)
        for i,v in enumerate(rowvals):
            c=row.cells[i]; c.text=str(v); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
            for p in c.paragraphs:
                for r in p.runs: r.font.size=Pt(7.8)
    return t

def callout(doc,title,text,fill=BLUE):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c=t.cell(0,0); shade(c,fill); margins(c,80,100)
    p=c.paragraphs[0]; r=p.add_run(title+' — '); r.bold=True; r.font.color.rgb=RGBColor.from_string(DARK); p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

def code(doc, text):
    t=doc.add_table(rows=1,cols=1); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    nosplit(t.rows[0])
    c=t.cell(0,0); shade(c,'F7F7F7'); margins(c,65,90)
    p=c.paragraphs[0]; r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(7.5)

def bullet(doc,text,level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text)

def num(doc,text):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text)

def build(out: Path=OUT):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.68); sec.right_margin=Inches(.68)
    styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.1); styles['Normal'].paragraph_format.space_after=Pt(4); styles['Normal'].paragraph_format.line_spacing=1.05
    for name,size,col in [('Title',24,'17365D'),('Heading 1',15,'17365D'),('Heading 2',11.5,'2F5597'),('Heading 3',10,'4472C4')]:
        styles[name].font.name='Aptos Display'; styles[name].font.size=Pt(size); styles[name].font.color.rgb=RGBColor.from_string(col)

    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('FULL-ALBUM-MAKER'); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string('4472C4')
    p=doc.add_paragraph(style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run('STEP 07\nVisual Property Binding & Core Beat Animations')
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Beat Animation Engine V2 — Planning & Implementation Contract'); r.bold=True; r.font.size=Pt(11)
    callout(doc,'Status STEP','STEP 07 mengikat signal kontinu STEP 06 ke properti visual yang netral terhadap renderer. Tidak ada UI editor baru dan tidak ada rumus beat yang disalin ke preview/FFmpeg secara terpisah.',GREEN)
    table(doc,['Item','Nilai'],[
        ('Repository','inoriko920-dev/Full-Album-Maker'),('Development branch','feature/beat-animation-engine-v2'),('Stable main baseline','584c94774e6197ecf60ecedb3bffd5a8797e7737 (v1.5.0)'),('STEP 06 verified head','ac1e2c4e123258e804cc32163ea894a92eace4d8'),('Tanggal','7 Oktober 2026 (WIB)'),('Target','Satu adapter signal→visual property untuk preview dan final renderer berikutnya')])
    doc.add_page_break()

    doc.add_heading('1. Tujuan STEP 07',level=1)
    doc.add_paragraph('STEP 06 sudah menjawab bagaimana nilai beat/bass/onset berkembang terhadap waktu. STEP 07 tidak boleh mengulang attack/decay. Tugasnya hanya menerjemahkan nilai signal 0..1 menjadi delta visual yang aman terhadap transform manual pengguna.')
    for x in ['Transform manual tetap source baseline.','Beat animation bersifat non-destructive dan kembali tepat ke baseline ketika signal=0.','Satu binding engine dipakai semua layer/consumer.','Preview dan renderer berikutnya menerima EffectiveVisualState yang sama.','Tidak ada random direction pada V1; seluruh hasil deterministik.']: bullet(doc,x)
    callout(doc,'Aturan paling penting','Tidak boleh menulis hasil pulse ke Layer.transform setiap frame. Layer.transform tetap data authoring; hasil beat hanya effective state sementara.',YELLOW)

    doc.add_heading('2. Ruang Lingkup',level=1)
    table(doc,['Masuk STEP 07','Tidak Masuk STEP 07'],[
        ('Visual binding contract','Panel UI animation'),('Property evaluator','FFmpeg expression integration'),('Pivot-safe scale application','Preview playback integration'),('Core preset binding sets','Particle system'),('Transform/opacity effective adapter','Genre preset system'),('Glow/zoom logical outputs','AI Agent animation command'),('Focused + regression tests','DROP/SECTION detection')])

    doc.add_heading('3. Input Contract dari STEP 06',level=1)
    table(doc,['Input','Makna'],[
        ('AnimationSignalEngine.value/sample','Nilai channel 0..1 pada project tick'),('beat','Pulse ringan'),('strong_beat','Punch utama'),('bass','Impact low-frequency'),('onset','Accent singkat'),('energy_up/down','Perubahan intensitas makro')])

    doc.add_page_break()
    doc.add_heading('4. Visual Property Vocabulary V1',level=1)
    table(doc,['Property','Neutral','Range output V1','Tujuan'],[
        ('SCALE_MULTIPLIER','1.0','0.75..1.35','Pulse elemen/cover'),('ZOOM_MULTIPLIER','1.0','1.0..1.25','Punch crop/camera adapter nanti'),('OPACITY_MULTIPLIER','1.0','0.0..1.25','Dip/fade relative terhadap opacity base'),('ROTATION_OFFSET_DEG','0°','-12..+12°','Nudge non-akumulatif'),('X_OFFSET_NORMALIZED','0','-0.15..+0.15','Motion kecil relatif canvas'),('Y_OFFSET_NORMALIZED','0','-0.15..+0.15','Motion kecil relatif canvas'),('GLOW_AMOUNT','0','0..1','Logical glow intensity untuk consumer berikutnya')])
    callout(doc,'Neutral-state guarantee','Jika seluruh signal 0, effective state wajib identik dengan baseline authoring: scale=1, zoom=1, opacity multiplier=1, rotation/x/y/glow offset=0.',GREEN)

    doc.add_heading('5. Kontrak Data Baru',level=1)
    code(doc,"""VISUAL_BINDING_ENGINE_VERSION = 'visual-bindings-v1'

VisualProperty enum
VisualBinding:
    channel
    property
    amount
    response_gamma
    min_signal

VisualBindingSet:
    binding_set_id
    bindings: tuple[VisualBinding, ...]

VisualPropertyState:
    scale_multiplier=1
    zoom_multiplier=1
    opacity_multiplier=1
    rotation_offset_deg=0
    x_offset_normalized=0
    y_offset_normalized=0
    glow_amount=0

EffectiveLayerVisual:
    x, y, width, height, rotation, opacity
    zoom_multiplier, glow_amount""")

    doc.add_heading('6. Transfer Function',level=1)
    code(doc,"""s = signal_engine.value(channel, tick)
if s < min_signal:
    response = 0
else:
    response = ((s-min_signal)/(1-min_signal)) ** response_gamma
contribution = amount * response""")
    bullet(doc,'response_gamma mengatur karakter respons properti, bukan attack/decay. Temporal envelope tetap milik STEP 06.')
    bullet(doc,'amount boleh signed untuk rotation/x/y/opacity dip; scale/zoom/glow preset V1 default positif.')

    doc.add_page_break()
    doc.add_heading('7. Property Combination Rules',level=1)
    table(doc,['Property','Kombinasi multi-binding','Clamp'],[
        ('Scale','1 + sum(delta)','0.75..1.35'),('Zoom','1 + sum(delta)','1.0..1.25'),('Opacity multiplier','1 + sum(delta)','0..1.25'),('Rotation','sum(offset)','-12..12°'),('X/Y offset','sum(offset)','-0.15..0.15'),('Glow','saturating add: 1-Π(1-c)','0..1')])
    callout(doc,'Kenapa bukan MAX untuk semua?','Binding berbeda dapat sengaja digabung, misalnya beat pulse kecil + bass pulse lebih besar. Hard clamp menjaga hasil tetap profesional dan tidak meledak.',BLUE)

    doc.add_heading('8. Pivot-Safe Scale Application',level=1)
    doc.add_paragraph('Scale visual harus mempertahankan pivot transform. Mengubah width/height tanpa x/y akan membuat cover tumbuh dari pojok.')
    code(doc,"""pivot_abs_x = base.x + base.pivot_x * base.width
pivot_abs_y = base.y + base.pivot_y * base.height
new_width  = base.width  * scale_multiplier
new_height = base.height * scale_multiplier
new_x = pivot_abs_x - base.pivot_x * new_width  + x_offset
new_y = pivot_abs_y - base.pivot_y * new_height + y_offset
new_rotation = base.rotation + rotation_offset
new_opacity = clamp(base_opacity * opacity_multiplier, 0..1)""")
    bullet(doc,'Final width/height tetap diklem ke batas renderer 0.02..3.0.')
    bullet(doc,'Final rotation tetap diklem -180..180 derajat.')
    bullet(doc,'Base Transform object tidak dimutasi.')

    doc.add_heading('9. Logical Zoom vs Element Scale',level=1)
    doc.add_paragraph('SCALE_MULTIPLIER mengubah geometri layer. ZOOM_MULTIPLIER sengaja dipisahkan sebagai logical crop/camera zoom agar background/video dapat zoom tanpa mengubah slot layout. STEP 07 hanya menghasilkan nilainya; integrasi renderer dilakukan pada STEP berikutnya.')

    doc.add_page_break()
    doc.add_heading('10. Enam Core Beat Animations V1',level=1)
    table(doc,['Preset','Binding','Karakter'],[
        ('Subtle Beat Pulse','beat → scale +0.025','Gerak ringan kontinu'),('Bass Pulse','bass → scale +0.060','Cover/visual menghentak bass'),('Strong Punch','strong_beat → scale +0.080; zoom +0.060; glow +0.35','Accent utama tanpa shake'),('Onset Flash','onset → glow +0.28','Transient sparkle/flash logical'),('Rotation Nudge','strong_beat → rotation +2.0°','Nudge singkat kembali ke base'),('Energy Breathe','energy_up → scale +0.030; glow +0.20','Build-up lembut')])
    callout(doc,'V1 philosophy','Sedikit animasi inti yang stabil lebih baik daripada puluhan efek yang masing-masing punya matematika berbeda. Setelah adapter stabil, jumlah preset dapat tumbuh tanpa mengubah engine.',GREEN)

    doc.add_heading('11. Preset Contract',level=1)
    code(doc,"""CoreBeatPreset enum:
    SUBTLE_BEAT_PULSE
    BASS_PULSE
    STRONG_PUNCH
    ONSET_FLASH
    ROTATION_NUDGE
    ENERGY_BREATHE

core_binding_set(preset) -> VisualBindingSet
combine_binding_sets(*sets) -> VisualBindingSet""")
    bullet(doc,'Preset hanyalah kumpulan binding; bukan branch kode khusus di renderer.')
    bullet(doc,'Pengguna nanti dapat mengubah amount/gamma melalui UI tanpa mengubah engine.')

    doc.add_heading('12. Layer Compatibility V1',level=1)
    table(doc,['Layer type','Scale','Zoom','Opacity','Rotation','Glow'],[
        ('song_cover','YA','YA','YA','YA','YA'),('vinyl','YA','YA','YA','YA','YA'),('background','YA','YA','YA','terbatas','YA'),('song_visual','YA','YA','YA','YA','YA'),('spectrum','YA','tidak V1','YA','YA','YA'),('text/song_title','YA','tidak V1','YA','YA','YA'),('progress/song_time','opsional','tidak','YA','opsional','tidak')])
    doc.add_paragraph('Compatibility table adalah policy helper untuk UI/consumer berikutnya; core evaluator tetap generic.')

    doc.add_page_break()
    doc.add_heading('13. Preview/Renderer Parity Contract',level=1)
    callout(doc,'Single math path','Preview dan renderer dilarang menghitung pulse sendiri. Keduanya harus meminta VisualPropertyState/EffectiveLayerVisual dari modul STEP 07 untuk tick yang sama.',YELLOW)
    for x in ['Same base transform + same signal program + same binding set + same tick => effective visual identik.','No wall-clock time. Semua berbasis project tick.','No random jitter di V1.','FFmpeg integration nanti boleh mengompilasi rumus, tetapi golden tests harus dibandingkan terhadap evaluator Python.']: bullet(doc,x)

    doc.add_heading('14. Seek, Scrub, dan Frame Rate',level=1)
    doc.add_paragraph('Binding evaluator bersifat random-access. FPS tidak masuk ke perhitungan nilai. Render 30fps dan preview 60fps yang bertemu pada project tick sama harus menghasilkan state yang sama.')

    doc.add_heading('15. Performance',level=1)
    table(doc,['Operasi','Target'],[
        ('Evaluate one binding set','O(B + signal lookups), B kecil'),('Signal sampling','Deduplicate channel; sample sekali per tick'),('Apply effective transform','O(1)'),('Preset combine','O(total bindings), dilakukan saat konfigurasi'),('Album length','Tidak memengaruhi ukuran binding set')])

    doc.add_heading('16. Serialization / Project Schema',level=1)
    doc.add_paragraph('STEP 07 belum mengubah ProjectDocument schema_version. BindingSet belum otomatis dipersist ke layer.animation. Persistence/UI schema akan dikunci pada STEP integrasi editor agar migration dapat dirancang sekaligus.')
    bullet(doc,'Tidak menulis runtime state ke Layer.transform.')
    bullet(doc,'Tidak mengubah revision/dirty state saat evaluate.')
    bullet(doc,'Factory preset dan binding objects immutable.')

    doc.add_page_break()
    doc.add_heading('17. Validation Rules',level=1)
    table(doc,['Field','Aturan'],[
        ('amount','finite; property-specific safe range'),('response_gamma','>0 dan finite'),('min_signal','0..0.95'),('channel','AnimationSignalChannel valid'),('property','VisualProperty valid'),('binding_set_id','non-empty'),('duplicate binding','boleh jika property/channel sama; combine deterministic'),('tick','diteruskan ke signal engine; negative ditolak')])

    doc.add_heading('18. Failure Semantics',level=1)
    table(doc,['Kondisi','Respons'],[
        ('Unknown preset','KeyError'),('Invalid binding','ValueError sebelum evaluate'),('Signal channel unavailable','KeyError / fail closed pada optional extension; core V1 reject'),('Empty binding set','Valid neutral state'),('Base opacity invalid','ValueError'),('Transform invalid','Validate before apply'),('Property overflow','Clamp safe; bukan exception runtime')])

    doc.add_heading('19. Module Plan',level=1)
    code(doc,"""src/full_album_maker/
    visual_binding_contract.py
    visual_binding_engine.py

tests/
    test_step07_visual_bindings.py

.github/workflows/
    beat-animation-step07.yml""")

    doc.add_heading('20. Test Matrix',level=1)
    tests=[
        'Binding validation valid','Reject invalid amount/gamma/min_signal','BindingSet signature deterministic','Signature changes on binding change','Empty set returns neutral state','Beat pulse scale exact','Bass pulse scale exact','Strong punch emits scale+zoom+glow','Onset flash glow exact','Rotation nudge returns to baseline','Energy breathe values','Signal below min_signal gated','response_gamma transfer correct','Multiple scale bindings sum then clamp','Glow saturating add bounded','Opacity multiplier bounded','Rotation clamp bounded','X/Y clamp bounded','Pivot center preserved during scale','Non-center pivot preserved','Base Transform not mutated','Base opacity not mutated','Effective width/height renderer-safe','Effective rotation renderer-safe','Same tick deterministic regardless binding order','Combined preset deterministic','Same signal sample not queried redundantly per channel','Random seek state equals repeated evaluation','Signal zero exactly restores baseline','Core preset IDs unique','Layer compatibility helper deterministic','STEP06 regression PASS','STEP05 regression PASS','STEP04 regression PASS']
    table(doc,['ID','Test'],[(f'{i:02d}',t) for i,t in enumerate(tests,1)])

    doc.add_page_break()
    doc.add_heading('21. CI Gate',level=1)
    for x in ['Windows + Ubuntu: STEP07 focused tests.','Windows + Ubuntu: STEP06, STEP05, STEP04 regression.','Import smoke visual_binding_contract + visual_binding_engine.','Tidak merge ke main.']: bullet(doc,x)

    doc.add_heading('22. Urutan Implementasi',level=1)
    for x in ['Commit generator/workflow STEP07 planning.','Generate dan commit DOCX STEP07 + update SOURCE_OF_TRUTH.','Implement visual_binding_contract.py.','Implement visual_binding_engine.py dan pivot-safe adapter.','Implement 6 core preset factories.','Tambah focused tests.','Tambah CI Windows/Ubuntu + regression.','Audit main tetap baseline.']: num(doc,x)

    doc.add_heading('23. Acceptance Gate',level=1)
    table(doc,['Gate','Kriteria PASS'],[
        ('G1 Source of Truth','STEP07 DOCX committed sebelum code'),('G2 Neutral','Signal=0 mengembalikan baseline persis'),('G3 Pivot','Scale menjaga pivot'),('G4 Bounds','Output property dalam safe limits'),('G5 Determinism','Urutan binding/input tidak mengubah hasil'),('G6 Parity Contract','Satu evaluator consumer-neutral tersedia'),('G7 No Mutation','Project/Transform base tidak dimutasi'),('G8 Regression','STEP04-06 PASS Windows+Ubuntu'),('G9 Stable Main','main tetap baseline v1.5.0')])
    callout(doc,'Gate keputusan','Semua G1-G9 wajib PASS sebelum STEP 07 ditutup.',GREEN)

    doc.add_heading('24. Definition of Done',level=1)
    for x in ['Visual property binding engine tersedia dan immutable.','7 property outputs V1 tersedia.','6 core beat-animation presets tersedia.','Pivot-safe effective transform tersedia.','Signal zero kembali tepat ke transform manual.','Tidak ada UI/FFmpeg logic duplikat.','Focused + regression tests PASS Windows/Ubuntu.']: bullet(doc,x)

    doc.add_heading('25. Handoff ke STEP 08',level=1)
    doc.add_paragraph('Setelah STEP 07 PASS, STEP 08 baru mengintegrasikan binding ke preview nyata dan jalur render/compile dengan parity tests. STEP 08 wajib memakai evaluator/adapter STEP 07, bukan menulis rumus beat baru di preview atau FFmpeg.')
    callout(doc,'Ringkasan final','STEP 06 menghasilkan signal. STEP 07 mengubah signal menjadi properti visual efektif. STEP 08 baru menampilkan dan merender properti tersebut.',BLUE)

    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run('Full-Album-Maker • Beat Animation Engine V2 • STEP 07 • 7 Oct 2026'); r.font.size=Pt(7); r.font.color.rgb=RGBColor(110,110,110)
    out.parent.mkdir(parents=True,exist_ok=True); doc.save(out)
    digest = hashlib.sha256(out.read_bytes()).hexdigest()
    text = SOT.read_text(encoding='utf-8') if SOT.exists() else '# Beat Animation Engine V2 — Source of Truth\n'
    entry = f"- `{out.name}` — SHA256 `{digest}`"
    lines = [line for line in text.splitlines() if out.name not in line]
    insert_at = next((i for i,line in enumerate(lines) if line.startswith('## Handoff rule')), len(lines))
    lines.insert(insert_at, entry); lines.insert(insert_at+1, '')
    SOT.write_text('\n'.join(lines).rstrip() + '\n', encoding='utf-8')
    print(f'Generated {out} sha256={digest}')

if __name__=='__main__': build()
