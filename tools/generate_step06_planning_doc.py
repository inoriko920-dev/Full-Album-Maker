from __future__ import annotations

import hashlib
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOC_ROOT = ROOT / "docs" / "beat-animation"
OUT = DOC_ROOT / "STEP06_ANIMATION_SIGNAL_ENVELOPE_ENGINE_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT = DOC_ROOT / "SOURCE_OF_TRUTH.md"

SECTIONS = [
("1. Masalah yang Diselesaikan", [
"STEP 05 memberi event diskrit, tetapi efek visual tidak boleh membaca event sebagai ON/OFF karena hasil akan patah, flicker, dan berbeda antar efek.",
"STEP 06 membuat satu Animation Signal / Envelope Engine sehingga cover pulse, glow, zoom, spectrum punch, dan efek lain memakai definisi attack, hold, decay, smoothing, intensity, dan overlap yang sama.",
]),
("2. Ruang Lingkup", [
"Masuk: envelope attack/hold/decay, routing per event type, overlap arbitration, easing, sensitivity, intensity, project-tick evaluator, range sampling.",
"Tidak masuk: UI animation inspector, visual effect final, particle renderer, Spectrum V2, preset genre, AI Agent binding, atau perubahan FFmpeg production renderer.",
]),
("3. Input Contract dari STEP 05", [
"Input utama adalah ProjectedMusicEvent: project_tick, source_tick, song_id, asset_id, event_type, strength, confidence, source.",
"Engine tidak mendefinisikan ulang beat. STRONG_BEAT/BASS_HIT/ENERGY_RISE/FALL tetap mengikuti STEP 05.",
]),
("4. Kontrak Data V1", [
"ANIMATION_SIGNAL_ENGINE_VERSION = animation-signals-v1.",
"SignalProfile menyimpan channel, event type, shape, attack_ms, hold_ms, decay_ms, strength_gamma, confidence_floor, amplitude, blend_mode, max_value.",
"AnimationSignalProgram immutable: duration_tick, settings_signature, channels, sorted trigger tuples.",
"AnimationSignalSample immutable dan berisi tick + channel values 0..1.",
]),
("5. Channel Signal V1", [
"beat <- BEAT; strong_beat <- STRONG_BEAT; bass <- BASS_HIT; onset <- ONSET; energy_up <- ENERGY_RISE; energy_down <- ENERGY_FALL.",
"Cross-channel event pada tick sama tidak di-dedupe; masing-masing tetap independen.",
]),
("6. Default Signal Profiles", [
"beat: attack 0ms, hold 0ms, decay 180ms, gamma 1.00, confidence floor 0.45, amplitude 0.65.",
"strong_beat: attack 12ms, hold 28ms, decay 320ms, gamma 0.85, floor 0.50, amplitude 1.00.",
"bass: attack 8ms, hold 20ms, decay 260ms, gamma 0.82, floor 0.35, amplitude 1.00.",
"onset: attack 0ms, hold 0ms, decay 120ms, gamma 1.10, floor 0.20, amplitude 0.60.",
"energy_up: attack 100ms, hold 180ms, decay 900ms, gamma 1.00, floor 0.40, amplitude 0.80.",
"energy_down: attack 80ms, hold 160ms, decay 760ms, gamma 1.00, floor 0.40, amplitude 0.75.",
]),
("7. Trigger Amplitude", [
"effective strength = clamp(event.strength * sensitivity, 0..1).",
"Jika confidence < confidence_floor maka trigger amplitude = 0.",
"strength_term = effective_strength ** strength_gamma.",
"confidence dipetakan dari confidence_floor..1 ke 0.35..1 agar MEDIUM event yang valid tidak dibuat terlalu lemah.",
"trigger amplitude = clamp(profile amplitude * strength_term * confidence_term, 0..max_value).",
"Intensity diterapkan setelah envelope dan output final selalu clamp 0..1.",
]),
("8. Envelope Shapes", [
"IMPULSE_DECAY: puncak pada event tick lalu turun menurut easing hingga 0.",
"ATTACK_HOLD_DECAY: optional pre-attack menuju event center, hold, lalu decay.",
"Easing V1: linear, smoothstep, ease_out_cubic; ease_out_expo optional dan tetap bounded.",
"BEAT/ONSET default tidak memakai pre-attack; STRONG_BEAT/BASS boleh pre-attack sangat pendek agar motion terasa natural.",
]),
("9. Overlap & Arbitration", [
"Raw SUM dilarang untuk output 0..1.",
"MAX menjadi default beat/strong_beat/bass.",
"SATURATING_ADD = 1-(1-a)(1-b) digunakan untuk onset/layered accent.",
"Duplikat channel pada project_tick sama menyimpan trigger amplitude terbesar.",
]),
("10. Sensitivity & Intensity", [
"Sensitivity 0..2 diterapkan sebelum gamma/envelope.",
"Intensity 0..2 diterapkan sesudah envelope.",
"Internal output tetap 0..1; UI nanti boleh menampilkan 0-200%.",
]),
("11. Evaluation API", [
"value(channel, tick) -> float.",
"sample(tick, channels=None) -> AnimationSignalSample.",
"sample_range(start_tick, end_tick, step_tick, channels=None) -> tuple samples.",
"Nilai pada tick T harus sama baik dicapai melalui seek acak maupun playback linear.",
]),
("12. Performance", [
"Trigger disimpan sebagai sorted tuple per channel.",
"Lookup single tick menggunakan bisect dan hanya mengevaluasi trigger yang envelope window-nya bisa aktif.",
"Tidak membuat buffer per-frame sepanjang album; album panjang tetap compact.",
]),
("13. Crossfade", [
"ProjectedMusicEvent sudah menangani posisi timeline. Pada crossfade dua lagu, event kedua lagu boleh overlap.",
"Overlap digabung per channel menurut blend rule. Audio gain/crossfade-volume-aware modulation ditunda; audio renderer tidak disentuh.",
]),
("14. Settings Signature", [
"SHA-256 mencakup engine version, profile, sensitivity, intensity, easing version.",
"Perubahan parameter yang memengaruhi output harus mengubah signature.",
"Program boleh memoized in-memory; tidak ada disk cache baru pada V1.",
]),
("15. Validation & Failure", [
"Timing >=0; gamma >0 finite; confidence floor 0..1; sensitivity/intensity 0..2.",
"Negative tick ditolak. Tick setelah duration menghasilkan 0. Unknown core channel ditolak.",
"Empty program valid dan semua channel bernilai 0.",
]),
("16. Thread Safety", [
"Program/profile/trigger/sample immutable.",
"Evaluator tidak memutasi ProjectDocument dan tidak memiliki worker thread sendiri.",
"Preview dan renderer boleh membaca program yang sama secara paralel.",
]),
("17. Module Plan", [
"src/full_album_maker/animation_signal_contract.py",
"src/full_album_maker/animation_signal_engine.py",
"tests/test_step06_animation_signals.py",
"Modul Qt-independent, librosa-independent, FFmpeg-independent.",
]),
("18. Test Matrix", [
"Minimal 32 focused tests: validation, signature, channel routing, confidence floor, sensitivity/intensity, attack/hold/decay boundary, monotonic decay, blend bounds, duplicate arbitration, seek determinism, sample_range parity, long sparse lookup, overlap dua lagu, immutability.",
"Regression wajib menjalankan STEP 05 dan STEP 04.",
]),
("19. CI Gate", [
"Windows dan Ubuntu menjalankan STEP06 focused + STEP05 + STEP04 regression.",
"Import smoke untuk contract/engine. Tidak merge ke main.",
]),
("20. Urutan Implementasi", [
"Commit DOCX STEP 06 terlebih dahulu dan update SOURCE_OF_TRUTH.",
"Implement contract/settings; program builder; envelope/easing; arbitration; random-access evaluator; sample_range; tests; CI.",
]),
("21. Acceptance Gate", [
"G1 DOCX committed sebelum code; G2 deterministic seek; G3 output 0..1; G4 smooth boundary tests; G5 no clipping; G6 no full-album frame buffer; G7 no ProjectDocument/UI/FFmpeg mutation; G8 STEP04/05 regression PASS Windows+Ubuntu; G9 main tetap baseline.",
]),
("22. Definition of Done", [
"Satu signal engine menjadi sumber envelope untuk Beat Engine V2.",
"6 channel V1 dapat dievaluasi pada arbitrary project tick.",
"Attack/hold/decay, easing, sensitivity, intensity, overlap terkontrak.",
"Focused dan regression tests PASS pada Windows dan Ubuntu.",
]),
("23. Handoff ke STEP 07", [
"STEP 07 baru mengikat signal ke properti visual dasar. STEP 07 dilarang mendefinisikan ulang beat atau smoothing; semua adapter wajib membaca AnimationSignalEngine.",
]),
]

def build() -> None:
    DOC_ROOT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.7); sec.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(9.5)
    for name, size, color in [("Title", 22, "17365D"), ("Heading 1", 15, "17365D"), ("Heading 2", 12, "2F5597")]:
        st = doc.styles[name]; st.font.name = "Aptos Display"; st.font.size = Pt(size); st.font.color.rgb = RGBColor.from_string(color)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("FULL-ALBUM-MAKER"); r.bold = True; r.font.color.rgb = RGBColor.from_string("4472C4")
    p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("STEP 06\nAnimation Signal / Envelope Engine")
    doc.add_paragraph("Branch: feature/beat-animation-engine-v2\nStable main baseline: 584c94774e6197ecf60ecedb3bffd5a8797e7737\nTanggal: 7 Oktober 2026 (WIB)")
    doc.add_heading("Arsitektur Inti", level=1)
    doc.add_paragraph("STEP 05 menjawab kapan event terjadi. STEP 06 menjawab bagaimana kekuatan event berkembang terhadap waktu. Visual adapter pada STEP berikutnya hanya menentukan properti apa yang digerakkan.")
    for title, paragraphs in SECTIONS:
        doc.add_heading(title, level=1)
        for text in paragraphs:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)
    footer = doc.sections[0].footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 06 • 7 Oct 2026")
    doc.save(OUT)

    digest = hashlib.sha256(OUT.read_bytes()).hexdigest()
    text = SOT.read_text(encoding="utf-8") if SOT.exists() else "# Beat Animation Engine V2 — Source of Truth\n"
    entry = f"- `{OUT.name}` — SHA256 `{digest}`"
    lines = [line for line in text.splitlines() if OUT.name not in line]
    insert_at = next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")), len(lines))
    lines.insert(insert_at, entry)
    lines.insert(insert_at + 1, "")
    SOT.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")
    print(f"Generated {OUT} sha256={digest}")

if __name__ == "__main__":
    build()
