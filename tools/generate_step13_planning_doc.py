from __future__ import annotations
import hashlib
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DOCS=ROOT/"docs"/"beat-animation"
OUT=DOCS/"STEP13_AI_MOTION_COMBINATION_PRESETS_PERFORMANCE_HARDENING_FULL_ALBUM_MAKER_2026-10-07.docx"
SOT=DOCS/"SOURCE_OF_TRUTH.md"
DARK="17365D"

def nosplit(row):
    p=row._tr.get_or_add_trPr()
    if not p.xpath("./w:cantSplit"):
        p.append(OxmlElement("w:cantSplit"))

def tbl(doc, headers, rows):
    t=doc.add_table(rows=1,cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; nosplit(t.rows[0])
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=h
        shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),DARK); c._tc.get_or_add_tcPr().append(shd)
        for p in c.paragraphs:
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for values in rows:
        row=t.add_row(); nosplit(row)
        for i,v in enumerate(values):
            row.cells[i].text=str(v)
            for p in row.cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(7.7)
    return t

def bullets(doc, values):
    for x in values:
        p=doc.add_paragraph(style="List Bullet"); p.add_run(x)

def code(doc, text):
    t=doc.add_table(rows=1,cols=1); t.style="Table Grid"; nosplit(t.rows[0])
    r=t.cell(0,0).paragraphs[0].add_run(text); r.font.name="Consolas"; r.font.size=Pt(7.2)

def build():
    doc=Document()
    sec=doc.sections[0]; sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(9)
    for name,size in [("Title",22),("Heading 1",14),("Heading 2",11)]:
        doc.styles[name].font.name="Aptos Display"; doc.styles[name].font.size=Pt(size); doc.styles[name].font.color.rgb=RGBColor.from_string(DARK)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("FULL-ALBUM-MAKER"); r.bold=True; r.font.color.rgb=RGBColor.from_string("4472C4")
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("STEP 13\nAI Motion Actions, Combination Presets & Performance Hardening")
    tbl(doc,["Item","Value"],[
        ("Development branch","feature/beat-animation-engine-v2"),
        ("Stable main","584c94774e6197ecf60ecedb3bffd5a8797e7737"),
        ("STEP12 verified head","599f15de6412cb05c9e18b984e813de3a5a5d130"),
        ("Scope","AI motion actions + 8 combo presets + command batching + long-album guards"),
    ])

    doc.add_heading("1. AI Motion Action Registry",1)
    tbl(doc,["Action","Arguments","Rule"],[
        ("set_beat_motion","layer_ids, motion_preset, motion_intensity","Preserve visual preset; validate capability"),
        ("adjust_motion_intensity","layer_ids, delta","Preserve motion preset; clamp 0..2"),
        ("clear_beat_motion","layer_ids","Remove motion only"),
        ("apply_beat_combo","layer_ids, combo_id","Write canonical visual+motion recipe"),
    ])
    bullets(doc,["All four actions require beat.write.","Preview Diff, atomic transaction and one Undo remain mandatory.","Ambiguous layer target must clarify instead of guessing."])

    doc.add_heading("2. Combination Preset Registry",1)
    tbl(doc,["Combo","Visual","Motion"],[
        ("club_impact","Club Punch 110%","Camera Shake 75%"),
        ("bass_rider","Bass Punch 105%","Bass Sway 85%"),
        ("neon_kick","Strong Glow 95%","Four-Way Kick 75%"),
        ("clean_bounce","Beat Zoom 80%","Beat Bounce 70%"),
        ("cinematic_spark","Cinematic Swell 80%","Spark Burst 65%"),
        ("remix_wobble","Club Punch 105%","Alternating Wobble 70%"),
        ("rock_shake","Bass Punch 95%","Camera Shake 65%"),
        ("ambient_breathe","Energy Glow 65%","Beat Bounce 25%"),
    ])
    bullets(doc,["Combo is convenience metadata, not a new engine.","Project persists canonical beat_v1 fields, not a required combo_id.","Inspector shows only combos whose visual and motion parts are supported by the layer."])

    doc.add_heading("3. Inspector, Context and NLU",1)
    code(doc,"Inspector: Combo / Beat Preset / Beat Intensity / Motion / Motion Intensity\nBeat Context: motion_catalog, combo_catalog, current_motion, motion_intensity, supported_motion_presets, supported_combos")
    tbl(doc,["Prompt","Action"],[
        ("tambahkan camera shake","set_beat_motion(camera_shake)"),
        ("pakai spark saat beat kuat","set_beat_motion(spark_burst)"),
        ("shake-nya lebih kuat","adjust_motion_intensity(+0.25)"),
        ("hapus motion","clear_beat_motion"),
        ("pakai Club Impact","apply_beat_combo(club_impact)"),
    ])
    bullets(doc,["Exact combo/motion aliases have priority over broad style aliases.","Gemini receives four new whitelisted function tools; no arbitrary project JSON."])

    doc.add_heading("4. Timestamp Command Batching",1)
    code(doc,"Before:\n12.500 scale@x width 420;\n12.500 scale@x height 420;\n12.500 eq@z brightness .03;\n\nAfter:\n12.500 scale@x width 420, scale@x height 420, eq@z brightness .03;")
    bullets(doc,["One row equals one timestamp batch.","Command order is deterministic and identical duplicate operations at the same tick are removed.","Timestamp, filter target, argument, sampling rate and evaluator state are unchanged."])

    doc.add_heading("5. Render Guards and Diagnostics",1)
    tbl(doc,["Guard","Limit"],[
        ("MAX_COMMAND_ROWS","250,000 timestamp batches"),
        ("MAX_COMMAND_OPS","1,500,000 operations"),
        ("MAX_COMMAND_FILE_BYTES","64 MiB"),
        ("Beat control","<=30 Hz; no automatic downgrade"),
        ("Spark","8 Hz inside active burst windows"),
        ("Beat layers","max 4"),
    ])
    code(doc,"BeatRenderControl diagnostics:\ncommand_rows / command_ops / command_bytes\nspark_command_rows / spark_command_ops")
    bullets(doc,["Guard failure happens before FFmpeg process start.","No silent lowering of visual quality to make a long album pass."])

    doc.add_heading("6. Long-Album Hardening",1)
    tbl(doc,["Fixture","Target"],[
        ("2h @120 BPM Strong Punch","<=216,000 timestamp batches"),
        ("2h + Camera Shake","rows sample-tick bounded; ops below guard"),
        ("2h + Spark Burst","rows bounded by 220ms strong-beat windows"),
        ("4 Beat layers","per-layer/global estimate before render"),
        ("1,000 random seeks","same tick -> identical state"),
    ])
    bullets(doc,["SparkBurstEngine uses EventPhaseEngine.active_triggers local window instead of scanning every strong beat on every tick.","BeatVisualRuntime state is evaluated once per layer/tick and reused for all commands."])

    doc.add_heading("7. Required Tests",1)
    tests=[
        "8 combos unique/valid/capability-safe","Inspector combo writes canonical payload","AI set/adjust/clear motion","AI apply combo + beat.write permission",
        "NLU camera shake/spark/stronger/clear/combo","Gemini tool schema includes four motion actions","AI plan card + one Undo",
        "CommandBatch one row per timestamp","deterministic op order and duplicate dedupe","row/op/byte guards",
        "2h estimator architecture","Spark local-window parity","real FFmpeg batched scale/glow","real FFmpeg overlay x/y","real FFmpeg Spark",
        "STEP04-12 regression Windows and Ubuntu",
    ]
    tbl(doc,["ID","Test"],[(str(i).zfill(2),x) for i,x in enumerate(tests,1)])

    doc.add_heading("8. Acceptance Gate",1)
    tbl(doc,["Gate","PASS Criteria"],[
        ("G1","DOCX committed before product code"),
        ("G2","8 combo recipes capability-safe"),
        ("G3","4 safe AI motion actions behind beat.write"),
        ("G4","NLU ambiguity fail-closed"),
        ("G5","Inspector/AI one Undo-safe transaction"),
        ("G6","same-timestamp commands batched"),
        ("G7","rows/ops/bytes preflight"),
        ("G8","no hidden sampling downgrade"),
        ("G9","Spark local lookup"),
        ("G10","STEP04-12 regression PASS Windows/Ubuntu"),
        ("G11","3 real-FFmpeg batching smokes PASS"),
        ("G12","stable main unchanged"),
    ])
    doc.add_heading("9. Handoff to STEP 14",1)
    doc.add_paragraph("STEP14 is final performance/reliability hardening: memory/cache profiling, cancellation, corrupt-state recovery, long-project stress, Windows portable dependency audit and release-candidate regression. No major feature addition after STEP13 unless a test proves a bug.")

    for s in doc.sections:
        p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=p.add_run("Full-Album-Maker • Beat Animation Engine V2 • STEP 13 • 7 Oct 2026"); r.font.size=Pt(7); r.font.color.rgb=RGBColor(105,105,105)

    DOCS.mkdir(parents=True,exist_ok=True); doc.save(OUT)
    digest=hashlib.sha256(OUT.read_bytes()).hexdigest()
    text=SOT.read_text(encoding="utf-8")
    bt=chr(96)
    entry="- "+bt+OUT.name+bt+" — SHA256 "+bt+digest+bt
    lines=[line for line in text.splitlines() if OUT.name not in line]
    idx=next((i for i,line in enumerate(lines) if line.startswith("## Handoff rule")),len(lines))
    lines.insert(idx,entry); lines.insert(idx+1,"")
    SOT.write_text("\n".join(lines).rstrip()+"\n",encoding="utf-8")
    print("Generated",OUT,"sha256",digest)

if __name__=="__main__":
    build()
